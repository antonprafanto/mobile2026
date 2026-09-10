"""
Generator Slide Pertemuan 01 - Versi Ramah Mahasiswa Pemula & Desain Minimalis Profesional
Menggunakan python-pptx dengan format 16:9 Widescreen.
Desain: Clean Minimalist Studio (Generous Whitespace, Tipografi Elegan, Card Terstruktur).
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# --- PALET WARNA CLEAN MINIMALIST PROFESIONAL ---
# Latar belakang bersih (Clean Studio Light-Slate atau Refined Dark Slate)
# Menggunakan tema Clean Minimalist Slate: Elegan, tidak silau, kontras tinggi, bebas warna neon norak.
COLOR_BG = RGBColor(248, 250, 252)          # #F8FAFC (Clean Light Surface)
COLOR_CARD = RGBColor(255, 255, 255)        # #FFFFFF (Pure White Card)
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 (Subtle Slate Border)
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A (Deep Slate Text)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B (Muted Secondary Text)
COLOR_ACCENT = RGBColor(2, 132, 199)        # #0284C7 (Refined Flutter Sky Blue)
COLOR_ACCENT_BG = RGBColor(240, 249, 255)   # #F0F9FF (Very Soft Blue Accent)
COLOR_SUCCESS = RGBColor(16, 185, 129)      # #10B981 (Subtle Green)
COLOR_WARNING = RGBColor(217, 119, 6)       # #D97706 (Warm Amber)
COLOR_CODE_BG = RGBColor(15, 23, 42)        # #0F172A (Deep Navy Dark for Code Box)
COLOR_CODE_TEXT = RGBColor(241, 245, 249)   # #F1F5F9 (Crisp Monospace Text)
COLOR_CODE_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 (Light Sky)

FONT_FAMILY = "Segoe UI"
FONT_CODE = "Consolas"

class MinimalistSlideBuilder:
    def __init__(self, course_name="IF3205 • Pemrograman Piranti Bergerak"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.course_name = course_name

    def _set_clean_bg(self, slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def _add_header(self, slide, tag, title, subtitle=None):
        # Subtle Tag Pill
        tb_tag = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.35))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.name = FONT_FAMILY
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT

        # Slide Main Title
        tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.8), Inches(11.5), Inches(0.65))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN

        if subtitle:
            tb_sub = slide.shapes.add_textbox(Inches(0.9), Inches(1.35), Inches(11.5), Inches(0.4))
            p_sub = tb_sub.text_frame.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(12.5)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # 1. Slide Cover Minimalis
    def add_cover(self, meeting_num, title, subtitle, presenter_info):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)

        # White Hero Card with Subtle Border
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.9), Inches(11.533), Inches(5.7)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"{self.course_name.upper()}  |  PERTEMUAN {meeting_num:02d}"
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT
        p0.space_after = Pt(20)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_MAIN
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(16)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(45)

        p3 = tf.add_paragraph()
        p3.text = presenter_info
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_TEXT_MAIN

    # 2. Slide Peta Jalan / Big Picture (Apa yang Akan Dibuat)
    def add_roadmap_slide(self, tag, title, subtitle, goals):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, tag, title, subtitle)

        # 4 Cards Milestone Sejajar
        num_cards = len(goals)
        card_w = Inches(2.65)
        gap = Inches(0.3)
        start_x = Inches(0.9)

        for i, goal in enumerate(goals):
            x = start_x + i * (card_w + gap)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), card_w, Inches(4.6)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_CARD
            card.line.color.rgb = COLOR_CARD_BORDER
            card.line.width = Pt(1)

            tb = slide.shapes.add_textbox(x + Inches(0.2), Inches(2.2), card_w - Inches(0.4), Inches(4.2))
            tf = tb.text_frame
            tf.word_wrap = True

            p_phase = tf.paragraphs[0]
            p_phase.text = f"FASE {i+1}"
            p_phase.font.name = FONT_FAMILY
            p_phase.font.size = Pt(10)
            p_phase.font.bold = True
            p_phase.font.color.rgb = COLOR_ACCENT
            p_phase.space_after = Pt(6)

            p_h = tf.add_paragraph()
            p_h.text = goal["title"]
            p_h.font.name = FONT_FAMILY
            p_h.font.size = Pt(15)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_TEXT_MAIN
            p_h.space_after = Pt(10)

            p_desc = tf.add_paragraph()
            p_desc.text = goal["desc"]
            p_desc.font.name = FONT_FAMILY
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = COLOR_TEXT_MUTED
            p_desc.space_after = Pt(14)

            p_out = tf.add_paragraph()
            p_out.text = f"Output: {goal['output']}"
            p_out.font.name = FONT_FAMILY
            p_out.font.size = Pt(11.5)
            p_out.font.bold = True
            p_out.font.color.rgb = COLOR_SUCCESS

    # 3. Slide Konsep Analogi (Flutter vs Dart)
    def add_analogy_slide(self, tag, title, subtitle, item_left, item_right):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, tag, title, subtitle)

        col_w = Inches(5.6)
        gap = Inches(0.33)

        # Left Card (Dart)
        card_l = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.0), col_w, Inches(4.6)
        )
        card_l.fill.solid()
        card_l.fill.fore_color.rgb = COLOR_CARD
        card_l.line.color.rgb = COLOR_CARD_BORDER
        card_l.line.width = Pt(1)

        tb_l = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), col_w - Inches(0.6), Inches(4.0))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        p_lt = tf_l.paragraphs[0]
        p_lt.text = item_left["role"].upper()
        p_lt.font.name = FONT_FAMILY
        p_lt.font.size = Pt(11)
        p_lt.font.bold = True
        p_lt.font.color.rgb = COLOR_ACCENT
        p_lt.space_after = Pt(8)

        p_lh = tf_l.add_paragraph()
        p_lh.text = item_left["name"]
        p_lh.font.name = FONT_FAMILY
        p_lh.font.size = Pt(22)
        p_lh.font.bold = True
        p_lh.font.color.rgb = COLOR_TEXT_MAIN
        p_lh.space_after = Pt(12)

        for pt in item_left["points"]:
            p = tf_l.add_paragraph()
            p.text = f"-  {pt}"
            p.font.name = FONT_FAMILY
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(8)

        # Right Card (Flutter)
        card_r = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9) + col_w + gap, Inches(2.0), col_w, Inches(4.6)
        )
        card_r.fill.solid()
        card_r.fill.fore_color.rgb = COLOR_CARD
        card_r.line.color.rgb = COLOR_CARD_BORDER
        card_r.line.width = Pt(1)

        tb_r = slide.shapes.add_textbox(Inches(1.2) + col_w + gap, Inches(2.3), col_w - Inches(0.6), Inches(4.0))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_rt = tf_r.paragraphs[0]
        p_rt.text = item_right["role"].upper()
        p_rt.font.name = FONT_FAMILY
        p_rt.font.size = Pt(11)
        p_rt.font.bold = True
        p_rt.font.color.rgb = COLOR_ACCENT
        p_rt.space_after = Pt(8)

        p_rh = tf_r.add_paragraph()
        p_rh.text = item_right["name"]
        p_rh.font.name = FONT_FAMILY
        p_rh.font.size = Pt(22)
        p_rh.font.bold = True
        p_rh.font.color.rgb = COLOR_TEXT_MAIN
        p_rh.space_after = Pt(12)

        for pt in item_right["points"]:
            p = tf_r.add_paragraph()
            p.text = f"-  {pt}"
            p.font.name = FONT_FAMILY
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(8)

    # 4. Slide Tiga Kolom Konseptual Minimalis
    def add_three_columns(self, tag, title, subtitle, col1, col2, col3):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, tag, title, subtitle)

        cols = [col1, col2, col3]
        col_w = Inches(3.64)
        gap = Inches(0.3)
        start_x = Inches(0.9)

        for i, col in enumerate(cols):
            x = start_x + i * (col_w + gap)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), col_w, Inches(4.6)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_CARD
            card.line.color.rgb = COLOR_CARD_BORDER
            card.line.width = Pt(1)

            tb = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.25), col_w - Inches(0.5), Inches(4.1))
            tf = tb.text_frame
            tf.word_wrap = True

            p_tag = tf.paragraphs[0]
            p_tag.text = col.get("tag", "").upper()
            p_tag.font.name = FONT_FAMILY
            p_tag.font.size = Pt(10)
            p_tag.font.bold = True
            p_tag.font.color.rgb = COLOR_ACCENT
            p_tag.space_after = Pt(6)

            p_h = tf.add_paragraph()
            p_h.text = col["title"]
            p_h.font.name = FONT_FAMILY
            p_h.font.size = Pt(17)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_TEXT_MAIN
            p_h.space_after = Pt(10)

            for pt in col["points"]:
                p = tf.add_paragraph()
                p.text = f"-  {pt}"
                p.font.name = FONT_FAMILY
                p.font.size = Pt(12.5)
                p.font.color.rgb = COLOR_TEXT_MUTED
                p.space_after = Pt(6)

    # 5. Slide Anatomi Kode Minimalis
    def add_code_anatomy(self, tag, title, subtitle, bullets, code_snippet, note=None):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, tag, title, subtitle)

        # Left Explanations Card
        card_l = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.0), Inches(4.8), Inches(4.6)
        )
        card_l.fill.solid()
        card_l.fill.fore_color.rgb = COLOR_CARD
        card_l.line.color.rgb = COLOR_CARD_BORDER
        card_l.line.width = Pt(1)

        tb_l = slide.shapes.add_textbox(Inches(1.15), Inches(2.25), Inches(4.3), Inches(4.1))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        p_lt = tf_l.paragraphs[0]
        p_lt.text = "ANATOMI KODE"
        p_lt.font.name = FONT_FAMILY
        p_lt.font.size = Pt(10)
        p_lt.font.bold = True
        p_lt.font.color.rgb = COLOR_ACCENT
        p_lt.space_after = Pt(10)

        for b in bullets:
            p = tf_l.add_paragraph()
            p.text = f"-  {b}"
            p.font.name = FONT_FAMILY
            p.font.size = Pt(12.5)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_after = Pt(10)

        if note:
            p_note = tf_l.add_paragraph()
            p_note.text = f"💡 {note}"
            p_note.font.name = FONT_FAMILY
            p_note.font.size = Pt(11.5)
            p_note.font.color.rgb = COLOR_WARNING

        # Right Dark Code Box
        card_r = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(2.0), Inches(6.533), Inches(4.6)
        )
        card_r.fill.solid()
        card_r.fill.fore_color.rgb = COLOR_CODE_BG
        card_r.line.color.rgb = COLOR_TEXT_MUTED
        card_r.line.width = Pt(0.5)

        tb_r = slide.shapes.add_textbox(Inches(6.2), Inches(2.2), Inches(6.0), Inches(4.2))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_rt = tf_r.paragraphs[0]
        p_rt.text = "main.dart"
        p_rt.font.name = FONT_CODE
        p_rt.font.size = Pt(11)
        p_rt.font.color.rgb = COLOR_CODE_ACCENT
        p_rt.space_after = Pt(10)

        p_code = tf_r.add_paragraph()
        p_code.text = code_snippet
        p_code.font.name = FONT_CODE
        p_code.font.size = Pt(11)
        p_code.font.color.rgb = COLOR_CODE_TEXT

    # 6. Slide Panduan Visual Hardware / Debugging Praktis
    def add_step_guide(self, tag, title, subtitle, steps):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, tag, title, subtitle)

        step_w = Inches(3.64)
        gap = Inches(0.3)
        start_x = Inches(0.9)

        for i, step in enumerate(steps):
            x = start_x + i * (step_w + gap)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), step_w, Inches(4.6)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_CARD
            card.line.color.rgb = COLOR_CARD_BORDER
            card.line.width = Pt(1)

            tb = slide.shapes.add_textbox(x + Inches(0.25), Inches(2.25), step_w - Inches(0.5), Inches(4.1))
            tf = tb.text_frame
            tf.word_wrap = True

            p_num = tf.paragraphs[0]
            p_num.text = f"LANGKAH {i+1:02d}"
            p_num.font.name = FONT_FAMILY
            p_num.font.size = Pt(11)
            p_num.font.bold = True
            p_num.font.color.rgb = COLOR_ACCENT
            p_num.space_after = Pt(6)

            p_h = tf.add_paragraph()
            p_h.text = step["title"]
            p_h.font.name = FONT_FAMILY
            p_h.font.size = Pt(16)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_TEXT_MAIN
            p_h.space_after = Pt(10)

            p_desc = tf.add_paragraph()
            p_desc.text = step["desc"]
            p_desc.font.name = FONT_FAMILY
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = COLOR_TEXT_MUTED
            p_desc.space_after = Pt(10)

            if "tip" in step:
                p_tip = tf.add_paragraph()
                p_tip.text = f"Tips: {step['tip']}"
                p_tip.font.name = FONT_FAMILY
                p_tip.font.size = Pt(11)
                p_tip.font.color.rgb = COLOR_WARNING

    # 7. Slide Lab Quest Ramah Pemula
    def add_lab_quest(self, meeting_num, title, time_minutes, goals, success_criteria):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_clean_bg(slide)
        self._add_header(slide, "PRAKTIKUM MANDIRI", f"Lab Quest Pertemuan {meeting_num:02d}: {title}")

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(11.533), Inches(4.9)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(Inches(1.3), Inches(2.1), Inches(10.7), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"⏱️ ESTIMASI: {time_minutes} MENIT  |  TARGET: REAL DEVICE SMARTPHONE / PC LAB"
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT
        p0.space_after = Pt(14)

        p1 = tf.add_paragraph()
        p1.text = "Langkah Pengerjaan Mandiri:"
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_MAIN
        p1.space_after = Pt(10)

        for g in goals:
            p = tf.add_paragraph()
            p.text = f"  [ ]  {g}"
            p.font.name = FONT_FAMILY
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(6)

        p_crit = tf.add_paragraph()
        p_crit.text = f"\nKriteria Sukses (Tunjukkan ke Asisten/Dosen):\n✓ {success_criteria}"
        p_crit.font.name = FONT_FAMILY
        p_crit.font.size = Pt(12)
        p_crit.font.bold = True
        p_crit.font.color.rgb = COLOR_SUCCESS

    def save(self, filepath):
        self.prs.save(filepath)
        print(f"[OK] Slide Minimalis Berhasil Disimpan: {filepath}")


# =====================================================================
# EKSEKUSI PENYUSUNAN 11 SLIDE PERTEMUAN 01 RAMAH PEMULA
# =====================================================================

output_dir = os.path.dirname(os.path.abspath(__file__))
deck1 = MinimalistSlideBuilder()

# Slide 1: Cover
deck1.add_cover(
    1,
    "Pengantar Pemrograman Piranti Bergerak\n& Ekosistem Flutter",
    "Fondasi, Arsitektur Aplikasi Modern, dan Persiapan Menjadi Mobile Engineer Profesional.",
    "Program Studi Teknik Informatika • Semester 5/6"
)

# Slide 2: The Big Picture (Roadmap Semester)
deck1.add_roadmap_slide(
    "GAMBARAN BESAR",
    "Apa yang Akan Kita Bangun Semester Ini?",
    "Peta jalan perkuliahan berbasis proyek (Project-Based Learning) dari nol hingga siap rilis portofolio.",
    [
        {
            "title": "Fondasi & UI",
            "desc": "Mempelajari bahasa Dart modern, menyusun hierarki widget, dan merancang antarmuka indah Material 3.",
            "output": "Katalog UI Responsif"
        },
        {
            "title": "State & Logic",
            "desc": "Memisahkan tampilan dari data menggunakan arsitektur Cubit/BLoC yang terprediksi.",
            "output": "Aplikasi Interaktif (UTS)"
        },
        {
            "title": "Data & Cloud",
            "desc": "Menghubungkan aplikasi ke RESTful API, database lokal luring (offline-first), dan Firebase BaaS.",
            "output": "Aplikasi Terkoneksi API"
        },
        {
            "title": "Sensor & Rilis",
            "desc": "Mengakses kamera, GPS, pengujian otomatis (Unit Testing), serta build APK release.",
            "output": "Produk Portofolio (UAS)"
        }
    ]
)

# Slide 3: Pertanyaan Mendasar Pemula (Flutter vs Dart)
deck1.add_analogy_slide(
    "KONSEP DASAR",
    "Memahami Perbedaan: Flutter vs Dart",
    "Banyak pemula mengira Flutter adalah bahasa pemrograman. Mari luruskan pemahaman ini.",
    {
        "role": "Bahasa Pemrograman (The Language)",
        "name": "Dart",
        "points": [
            "Dibuat oleh Google khusus untuk aplikasi client modern.",
            "Mendukung OOP murni (Class, Object, Interface).",
            "Memiliki fitur Sound Null Safety (mencegah bug NullPointer).",
            "Analogi: Dart adalah 'Bahan Baku & Logika Otak' aplikasi Anda."
        ]
    },
    {
        "role": "Kerangka Kerja UI (The Toolkit / Framework)",
        "name": "Flutter",
        "points": [
            "Kumpulan komponen visual (Button, Text, AppBar, Slider).",
            "Menyediakan mesin perender grafis ke layar HP.",
            "Satu kode jalan di Android, iOS, Web, dan Desktop.",
            "Analogi: Flutter adalah 'Kotak Perkakas Lego' yang siap disusun."
        ]
    }
)

# Slide 4: Mengapa Aplikasi Mobile & Peluang Kariernya?
deck1.add_three_columns(
    "RELEVANSI INDUSTRI",
    "Mengapa Mempelajari Mobile App Development?",
    "Peluang dan fakta kebutuhan industri perangkat lunak modern saat ini.",
    {
        "tag": "PASAR GLOBAL",
        "title": "Mobile-First Era",
        "points": [
            ">70% trafik internet global berasal dari smartphone.",
            "Pengguna menghabiskan rata-rata 4-5 jam sehari di aplikasi mobile.",
            "Setiap bisnis rintisan (startup) wajib memiliki mobile presence."
        ]
    },
    {
        "tag": "EFISIENSI WAKTU",
        "title": "Single Codebase",
        "points": [
            "Dulu: Harus merekrut tim Kotlin + tim Swift terpisah.",
            "Sekarang: 1 tim Flutter menghasilkan aplikasi Android & iOS sekaligus.",
            "Memangkas waktu dan biaya pengembangan hingga 50%."
        ]
    },
    {
        "tag": "KARIER MAHASISWA",
        "title": "Peluang Kerja Luas",
        "points": [
            "Permintaan Mobile Engineer konsisten tinggi di perbankan & startup.",
            "Portofolio aplikasi nyata di GitHub sangat memikat recruiter.",
            "Dapat menghasilkan produk mandiri (Indie Hacker / SaaS)."
        ]
    }
)

# Slide 5: Tiga Pendekatan Membangun Aplikasi Mobile
deck1.add_three_columns(
    "KOMPARASI TEKNOLOGI",
    "Tiga Cara Membangun Aplikasi Mobile",
    "Memahami posisi Flutter di antara berbagai pilihan teknologi pengembangan aplikasi.",
    {
        "tag": "PENDEKATAN 1",
        "title": "Native Murni",
        "points": [
            "Bahasa: Kotlin (Android) & Swift (iOS).",
            "Kelebihan: Akses fitur hardware paling mutakhir.",
            "Kekurangan: Codebase terpisah, biaya ganda, rilis fitur rawan terlambat satu sama lain."
        ]
    },
    {
        "tag": "PENDEKATAN 2",
        "title": "Hybrid WebView",
        "points": [
            "Bahasa: HTML, CSS, JavaScript (Cordova/Ionic).",
            "Kelebihan: Mudah bagi web developer.",
            "Kekurangan: Terasa seperti membuka website di dalam browser, performa patah-patah."
        ]
    },
    {
        "tag": "PENDEKATAN 3 (FLUTTER)",
        "title": "Canvas Engine",
        "points": [
            "Bahasa: Dart.",
            "Cara Kerja: Menggambar piksel langsung ke layar GPU.",
            "Kelebihan: Performa 60-120 fps, tampilan identik di semua HP, satu basis kode."
        ]
    }
)

# Slide 6: Cara Kerja Flutter (Analogi Ramah Pemula)
deck1.add_three_columns(
    "CARA KERJA FLUTTER",
    "Bagaimana Flutter Menampilkan Layar?",
    "Flutter tidak meminjam tombol bawaan Android/iOS, melainkan menggambarnya sendiri.",
    {
        "tag": "ANALOGI GAME ENGINE",
        "title": "Seperti Game Unity",
        "points": [
            "Game mobile menggambar karakter langsung ke layar kanvas tanpa tombol OS.",
            "Flutter menerapkan filosofi yang sama untuk aplikasi bisnis: tombol dan teks digambar mandiri oleh Impeller Engine."
        ]
    },
    {
        "tag": "KONSISTENSI PIKSEL",
        "title": "Bebas Masalah Fragmentasi",
        "points": [
            "Di Android native, tampilan bisa berbeda antara HP Samsung, Xiaomi, dan Google Pixel.",
            "Di Flutter: Tampilan aplikasi Anda dijamin 100% konsisten di semua merek HP."
        ]
    },
    {
        "tag": "PERFORMA TINGGI",
        "title": "Halus Tanpa Jeda",
        "points": [
            "Animasi transisi berjalan mulus pada 60 fps hingga 120 fps.",
            "Tidak ada jembatan penerjemah JavaScript (bebas jembatan penghambat waktu)."
        ]
    }
)

# Slide 7: Senjata Andalan: Stateful Hot Reload
deck1.add_three_columns(
    "PRODUKTIVITAS DEVELOPER",
    "Fitur 'Sihir' Flutter: Stateful Hot Reload",
    "Mengapa belajar Flutter sangat menyenangkan bagi mahasiswa pemula?",
    {
        "tag": "DULU (TRADISIONAL)",
        "title": "Kompilasi Lambat",
        "points": [
            "Ubah 1 warna tombol.",
            "Kompilasi ulang Gradle selama 2–3 menit.",
            "Aplikasi restart dari halaman login.",
            "Sangat membuang waktu dan menguji kesabaran."
        ]
    },
    {
        "tag": "SEKARANG (FLUTTER)",
        "title": "Hot Reload (<1 Detik)",
        "points": [
            "Ubah kode di VS Code.",
            "Tekan tombol 'Save' (Ctrl+S).",
            "Layar smartphone langsung berubah seketika dalam hitungan milidetik!",
            "State data tidak hilang."
        ]
    },
    {
        "tag": "MANFAAT BELAJAR",
        "title": "Eksplorasi Cepat",
        "points": [
            "Mahasiswa bisa langsung melihat hasil eksperimen UI tanpa menunggu.",
            "Belajar tata letak layout menjadi sangat intuitif dan menyenangkan."
        ]
    }
)

# Slide 8: Anatomi Kode Pertama Kita
deck1.add_code_anatomy(
    "BEDAH KODE PERDANA",
    "Membaca Berkas 'main.dart' Pertama Anda",
    "Jangan takut dengan baris kode! Mari pahami struktur dasarnya baris demi baris.",
    [
        "void main(): Titik awal gerbang eksekusi program dimulai.",
        "runApp(): Perintah untuk menempelkan widget ke layar smartphone.",
        "MaterialApp: Komponen pembungkus tema aplikasi standar Material Design.",
        "Scaffold: Struktur dasar halaman (memiliki AppBar, Body, dan FloatingActionButton).",
        "Center & Text: Widget untuk meletakkan tulisan tepat di tengah layar."
    ],
    """import 'package:flutter/material.dart';

// 1. Titik masuk aplikasi
void main() {
  runApp(const HaloApp());
}

// 2. Widget utama aplikasi kita
class HaloApp extends StatelessWidget {
  const HaloApp({super.key});

  @override
  Widget build(BuildContext context) {
    return const MaterialApp(
      home: Scaffold(
        body: Center(
          child: Text(
            'Halo Informatika 2026!',
            style: TextStyle(fontSize: 20),
          ),
        ),
      ),
    );
  }
}""",
    "Setiap elemen visual di Flutter adalah WIDGET yang disusun bertingkat (bercabang)!"
)

# Slide 9: Solusi Komputer Lab RAM 8GB (Real Device & Scrcpy)
deck1.add_step_guide(
    "PANDUAN PRAKTIKUM LAB",
    "Cara Aman Debugging Tanpa Bikin Komputer Lab Hang",
    "Komputer lab atau laptop Anda RAM 8GB? Jangan gunakan Android Emulator! Gunakan Real Device.",
    [
        {
            "title": "Aktifkan USB Debugging di HP",
            "desc": "Buka Settings -> About Phone -> Ketuk 'Build Number' sebanyak 7 kali berturut-turut hingga Developer Options aktif.",
            "tip": "Centang 'Selalu izinkan dari komputer ini' saat dialog muncul di layar HP."
        },
        {
            "title": "Gunakan Kabel Data Berkualitas",
            "desc": "Sambungkan smartphone ke komputer menggunakan kabel data USB (pastikan bukan sekadar kabel charger abal-abal).",
            "tip": "Ketik 'adb devices' di terminal untuk memverifikasi HP terdeteksi."
        },
        {
            "title": "Jalankan Scrcpy (Screen Copy)",
            "desc": "Ketik 'scrcpy' di terminal. Layar HP Anda akan muncul di monitor PC. Sangat ringan (RAM < 80MB)!",
            "tip": "Anda bisa mengklik layar HP langsung menggunakan mouse komputer."
        }
    ]
)

# Slide 10: Lab Quest Tantangan Praktikum Pertemuan 1
deck1.add_lab_quest(
    1,
    "Hello Flutter & The First Hot Reload",
    45,
    [
        "Jalankan 'flutter doctor' di terminal dan pastikan tidak ada tanda silang merah pada Flutter & Android SDK.",
        "Hubungkan smartphone Android Anda dan pastikan terdeteksi pada daftar perangkat (ketik 'flutter devices').",
        "Buat proyek baru dengan nama: 'flutter create tugas_pertemuan_01'.",
        "Buka proyek di VS Code, lalu jalankan aplikasi ke perangkat dengan perintah 'flutter run'.",
        "Modifikasi teks pada layar menjadi: 'Nama Lengkap Anda - NIM' dan ubah warna AppBar menjadi warna favorit Anda.",
        "Tekan tombol 'r' di terminal dan saksikan Hot Reload mengubah layar HP Anda dalam sekejap!"
    ],
    "Tunjukkan layar smartphone fisik yang menampilkan Nama & NIM Anda kepada Dosen atau Asisten Lab untuk mendapatkan nilai modul 1."
)

# Slide 11: Persiapan untuk Pertemuan Berikutnya
deck1.add_three_columns(
    "LANGKAH SELANJUTNYA",
    "Apa yang Perlu Dipersiapkan untuk Pertemuan 2?",
    "Minggu depan kita akan menyelami bahasa Dart: logika dan pemrograman asynchronous.",
    {
        "tag": "MATERI MINGGU DEPAN",
        "title": "Deep Dive Dart",
        "points": [
            "Logika pemrograman modern Dart.",
            "Sound Null Safety (bebas bug null).",
            "Konsep Asynchronous: Future, async, await, dan Stream data.",
            "Object-Oriented Programming di Dart."
        ]
    },
    {
        "tag": "TUGAS MANDIRI",
        "title": "Instalasi Mandiri",
        "points": [
            "Bagi yang membawa laptop pribadi, pastikan Flutter SDK dan VS Code sudah terpasang.",
            "Jika ada kendala instalasi, hubungi Asisten Lab di grup komunikasi kelas."
        ]
    },
    {
        "tag": "PESAN HARI INI",
        "title": "Mindset Belajar",
        "points": [
            "Jangan menghafal semua nama widget.",
            "Pahami pola dan hierarkinya.",
            "Eksperimen langsung adalah kunci tercepat menguasai Flutter."
        ]
    }
)

# Simpan berkas PPTX yang telah diaudit dan diperbarui
output_path = os.path.join(output_dir, "PERTEMUAN_01_Pengantar_Flutter_dan_Ekosistem.pptx")
deck1.save(output_path)
