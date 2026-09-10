"""
Generator Slide Pertemuan 06 (Edisi Masterpiece 18 Slide Lengkap & Super Ramah Awam)
Mata Kuliah: Pemrograman Piranti Bergerak (Flutter & Dart 2026)
Topik: State Management Reaktif (Cubit & BLoC Pattern Standar Industri 2026)
Standar Desain: Neo-Brutalism Murni (designprompts.dev/neo-brutalism)
- ZERO BORDER RADIUS: Sudut 90 derajat siku murni tanpa kelengkungan (MSO_SHAPE.RECTANGLE).
- BOLD SOLID BLACK BORDERS: Garis hitam tegas 2.5 Pt (#000000).
- HARD OFFSET BLOCK SHADOWS: Bayangan blok hitam solid +0.08 inch di belakang kartu.
- PALET NEO-BRUTALISM: Warm Cream (#FAF8F5), Canary Yellow (#FFE600), Electric Cyan (#38BDF8), Mint (#4ADE80), Coral (#FB7185), Lavender (#C084FC).
- FITUR UTAMA: 100% SLIDE PEMBAHASAN DISERTAI KODE NYATA FLUTTER, ANALOGI MEMBUMI, DAN TOMBOL TAUTAN GITHUB.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

COLOR_BG = RGBColor(250, 248, 245)        # #FAF8F5 (Warm Cream Canvas)
COLOR_BLACK = RGBColor(0, 0, 0)           # #000000 (Pure Black Border & Hard Shadows)
COLOR_WHITE = RGBColor(255, 255, 255)     # #FFFFFF (Pure White Card)
COLOR_YELLOW = RGBColor(255, 230, 0)      # #FFE600 (Canary Yellow)
COLOR_CYAN = RGBColor(56, 189, 248)       # #38BDF8 (Electric Cyan)
COLOR_MINT = RGBColor(74, 222, 128)       # #4ADE80 (Neo Mint / Green)
COLOR_CORAL = RGBColor(251, 113, 133)     # #FB7185 (Neo Coral / Pink)
COLOR_PURPLE = RGBColor(192, 132, 252)    # #C084FC (Lavender Purple)
COLOR_CODE_BG = RGBColor(18, 18, 18)      # #121212 (Deep Black Code Box)
COLOR_CODE_TEXT = RGBColor(245, 245, 245) # #F5F5F5 (Crisp Monospace)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

SHADOW_OFFSET = Inches(0.08)

GITHUB_BASE_URL = "https://github.com/antonprafanto/mobile2026/blob/main/07_contoh_kode_program/pertemuan_06_state_management_cubit"

class NeoBrutalistDeckBuilder:
    def __init__(self, course_name="IF3205 • PEMROGRAMAN PIRANTI BERGERAK"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.course_name = course_name

    def _set_canvas_bg(self, slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def _add_neobrutal_card(self, slide, x, y, w, h, fill_color=COLOR_WHITE, has_shadow=True, border_width=Pt(2.5)):
        if has_shadow:
            shadow = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, x + SHADOW_OFFSET, y + SHADOW_OFFSET, w, h
            )
            shadow.fill.solid()
            shadow.fill.fore_color.rgb = COLOR_BLACK
            shadow.line.fill.background()

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, y, w, h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = COLOR_BLACK
        card.line.width = border_width
        return card

    def _add_header(self, slide, tag_text, title_text, tag_color=COLOR_YELLOW):
        tb_meta = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), Inches(11.533), Inches(0.35))
        p_meta = tb_meta.text_frame.paragraphs[0]
        p_meta.text = self.course_name
        p_meta.font.name = FONT_HEADING
        p_meta.font.size = Pt(9)
        p_meta.font.bold = True
        p_meta.font.color.rgb = RGBColor(100, 100, 100)

        # Tag Badge
        tag = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(0.72), Inches(3.2), Inches(0.42))
        tag.fill.solid()
        tag.fill.fore_color.rgb = tag_color
        tag.line.color.rgb = COLOR_BLACK
        tag.line.width = Pt(2.0)

        tb_tag = slide.shapes.add_textbox(Inches(0.95), Inches(0.72), Inches(3.1), Inches(0.42))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag_text
        p_tag.font.name = FONT_HEADING
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_BLACK

        # Slide Title
        tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(11.533), Inches(0.55))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_BLACK

    # 1. Slide Cover
    def add_cover(self, meeting_num, title, subtitle, author):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)

        self._add_neobrutal_card(slide, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5), fill_color=COLOR_WHITE)

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(0.6))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_YELLOW
        top_bar.line.color.rgb = COLOR_BLACK
        top_bar.line.width = Pt(2.5)

        tb_tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.1), Inches(10.3), Inches(0.4))
        p_tb = tb_tb.text_frame.paragraphs[0]
        p_tb.text = f"{self.course_name}  ■  MODUL AJAR STANDAR 2026"
        p_tb.font.name = FONT_HEADING
        p_tb.font.size = Pt(11)
        p_tb.font.bold = True
        p_tb.font.color.rgb = COLOR_BLACK

        badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.6), Inches(2.0), Inches(2.8), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_CYAN
        badge.line.color.rgb = COLOR_BLACK
        badge.line.width = Pt(2.0)

        tb_badge = slide.shapes.add_textbox(Inches(1.6), Inches(2.05), Inches(2.8), Inches(0.5))
        p_badge = tb_badge.text_frame.paragraphs[0]
        p_badge.text = f"PERTEMUAN {meeting_num:02d}"
        p_badge.font.name = FONT_HEADING
        p_badge.font.size = Pt(14)
        p_badge.font.bold = True
        p_badge.font.color.rgb = COLOR_BLACK

        tb_main = slide.shapes.add_textbox(Inches(1.6), Inches(2.7), Inches(10.0), Inches(1.8))
        tf_main = tb_main.text_frame
        tf_main.word_wrap = True

        p_t = tf_main.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BLACK

        p_s = tf_main.add_paragraph()
        p_s.text = subtitle
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(12.5)
        p_s.font.color.rgb = RGBColor(60, 60, 60)
        p_s.space_before = Pt(10)

        tb_auth = slide.shapes.add_textbox(Inches(1.6), Inches(5.4), Inches(10.0), Inches(0.5))
        p_a = tb_auth.text_frame.paragraphs[0]
        p_a.text = f"👨‍🏫 Pengampu: {author}  •  Teknik Informatika"
        p_a.font.name = FONT_HEADING
        p_a.font.size = Pt(11)
        p_a.font.bold = True
        p_a.font.color.rgb = COLOR_BLACK

    # 2. Slide Konsep dengan Kode
    def add_concept_with_code(self, tag, title, bullets, code_snippet, filename="example.dart", tip=None, tag_color=COLOR_CYAN, full_code_file=None):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, tag, title, tag_color=tag_color)

        # Left White Card (4.8 Inches)
        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(4.8), Inches(5.1), fill_color=COLOR_WHITE)

        h_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(4.8), Inches(0.48))
        h_bar.fill.solid()
        h_bar.fill.fore_color.rgb = tag_color
        h_bar.line.color.rgb = COLOR_BLACK
        h_bar.line.width = Pt(2.5)

        tb_hb = slide.shapes.add_textbox(Inches(1.1), Inches(1.86), Inches(4.4), Inches(0.38))
        p_hb = tb_hb.text_frame.paragraphs[0]
        p_hb.text = "💡 KONSEP INTI & ANALOGI AWAM"
        p_hb.font.name = FONT_HEADING
        p_hb.font.size = Pt(10.5)
        p_hb.font.bold = True
        p_hb.font.color.rgb = COLOR_BLACK

        content_h = Inches(3.1) if tip else Inches(4.3)
        tb_l = slide.shapes.add_textbox(Inches(1.1), Inches(2.36), Inches(4.4), content_h)
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        for i, b in enumerate(bullets):
            p = tf_l.paragraphs[0] if i == 0 else tf_l.add_paragraph()
            p.text = f"■  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(9.8)
            p.font.color.rgb = RGBColor(20, 20, 20)
            p.space_after = Pt(6)

        if tip:
            tip_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.05), Inches(5.52), Inches(4.5), Inches(1.15))
            tip_box.fill.solid()
            tip_box.fill.fore_color.rgb = COLOR_YELLOW
            tip_box.line.color.rgb = COLOR_BLACK
            tip_box.line.width = Pt(2.0)

            tb_tip = slide.shapes.add_textbox(Inches(1.1), Inches(5.55), Inches(4.4), Inches(1.05))
            tf_tip = tb_tip.text_frame
            tf_tip.word_wrap = True
            p_tip = tf_tip.paragraphs[0]
            p_tip.text = f"💡 TIPS PRAKTIKUM & PEDOMAN:\n{tip}"
            p_tip.font.name = FONT_BODY
            p_tip.font.size = Pt(9.2)
            p_tip.font.bold = True
            p_tip.font.color.rgb = COLOR_BLACK

        # Right Sharp Black Code Box (6.433 Inches)
        self._add_neobrutal_card(slide, Inches(6.0), Inches(1.8), Inches(6.433), Inches(5.1), fill_color=COLOR_CODE_BG)

        c_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), Inches(1.8), Inches(6.433), Inches(0.48))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = tag_color
        c_bar.line.color.rgb = COLOR_BLACK
        c_bar.line.width = Pt(2.5)

        tb_cb = slide.shapes.add_textbox(Inches(6.2), Inches(1.86), Inches(6.0), Inches(0.38))
        p_cb = tb_cb.text_frame.paragraphs[0]
        p_cb.text = f"{filename}  |  KODE IMPLEMENTASI"
        p_cb.font.name = FONT_CODE
        p_cb.font.size = Pt(10.5)
        p_cb.font.bold = True
        p_cb.font.color.rgb = COLOR_BLACK

        code_box_h = Inches(3.68) if full_code_file else Inches(4.35)
        tb_r = slide.shapes.add_textbox(Inches(6.2), Inches(2.32), Inches(6.0), code_box_h)
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_code = tf_r.paragraphs[0]
        p_code.text = code_snippet
        p_code.font.name = FONT_CODE
        p_code.font.size = Pt(9.1)
        p_code.font.color.rgb = COLOR_CODE_TEXT

        if full_code_file:
            btn_x = Inches(6.15)
            btn_y = Inches(6.16)
            btn_w = Inches(6.133)
            btn_h = Inches(0.60)

            btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.04), btn_y + Inches(0.04), btn_w, btn_h)
            btn_shadow.fill.solid()
            btn_shadow.fill.fore_color.rgb = COLOR_BLACK
            btn_shadow.line.fill.background()

            btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
            btn.fill.solid()
            btn.fill.fore_color.rgb = COLOR_YELLOW
            btn.line.color.rgb = COLOR_BLACK
            btn.line.width = Pt(2.0)

            tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.08), btn_y + Inches(0.03), btn_w - Inches(0.16), btn_h - Inches(0.06))
            tf_btn = tb_btn.text_frame
            tf_btn.word_wrap = True

            p_btn = tf_btn.paragraphs[0]
            run_btn = p_btn.add_run()
            run_btn.text = f"▶ BUKA KODE LENGKAP DI GITHUB ({full_code_file})"
            run_btn.font.name = FONT_HEADING
            run_btn.font.size = Pt(9.6)
            run_btn.font.bold = True
            run_btn.font.color.rgb = COLOR_BLACK
            run_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

            p_btn_sub = tf_btn.add_paragraph()
            p_btn_sub.text = "💡 Buka di browser / salin source code: github.com/antonprafanto/mobile2026"
            p_btn_sub.font.name = FONT_BODY
            p_btn_sub.font.size = Pt(7.6)
            p_btn_sub.font.color.rgb = RGBColor(60, 60, 60)

    # 3. Slide Lab Quest
    def add_lab_quest(self, meeting_num, title, time_minutes, goals, success_criteria, full_code_file="slide_18_lab_quest_cart_cubit.dart"):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, "LAB QUEST MANDIRI", f"Pertemuan {meeting_num:02d}: {title}", tag_color=COLOR_YELLOW)

        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(11.533), Inches(5.1), fill_color=COLOR_WHITE)

        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(11.533), Inches(0.52))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_YELLOW
        banner.line.color.rgb = COLOR_BLACK
        banner.line.width = Pt(2.5)

        tb_b = slide.shapes.add_textbox(Inches(1.2), Inches(1.85), Inches(10.9), Inches(0.45))
        p_b = tb_b.text_frame.paragraphs[0]
        p_b.text = f"⏱️ ALOKASI: {time_minutes} MENIT  ■  TARGET UJI: BROWSER CHROME ATAU SMARTPHONE FISIK"
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(11.5)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_BLACK

        tb = slide.shapes.add_textbox(Inches(1.3), Inches(2.40), Inches(10.7), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_g = tf.paragraphs[0]
        p_g.text = "LANGKAH KERJA & SPESIFIKASI TUGAS:"
        p_g.font.name = FONT_HEADING
        p_g.font.size = Pt(12)
        p_g.font.bold = True
        p_g.font.color.rgb = COLOR_BLACK
        p_g.space_after = Pt(6)

        for g in goals:
            p = tf.add_paragraph()
            p.text = f"✔  {g}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.4)
            p.font.color.rgb = RGBColor(30, 30, 30)
            p.space_after = Pt(4)

        if full_code_file:
            btn_x = Inches(1.3)
            btn_y = Inches(5.8)
            btn_w = Inches(10.7)
            btn_h = Inches(0.85)

            btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.06), btn_y + Inches(0.06), btn_w, btn_h)
            btn_shadow.fill.solid()
            btn_shadow.fill.fore_color.rgb = COLOR_BLACK
            btn_shadow.line.fill.background()

            btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
            btn.fill.solid()
            btn.fill.fore_color.rgb = COLOR_YELLOW
            btn.line.color.rgb = COLOR_BLACK
            btn.line.width = Pt(2.0)

            tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.12), btn_y + Inches(0.03), btn_w - Inches(0.24), btn_h - Inches(0.06))
            tf_btn = tb_btn.text_frame
            tf_btn.word_wrap = True

            p_btn = tf_btn.paragraphs[0]
            run_btn = p_btn.add_run()
            run_btn.text = f"▶ BUKA KODE LENGKAP LAB QUEST DI GITHUB ({full_code_file})"
            run_btn.font.name = FONT_HEADING
            run_btn.font.size = Pt(11)
            run_btn.font.bold = True
            run_btn.font.color.rgb = COLOR_BLACK
            run_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

            p_btn_sub = tf_btn.add_paragraph()
            p_btn_sub.text = "💡 Buka di browser: Solusi lengkap Mini E-Commerce Cart & Wishlist dengan Cubit di GitHub"
            p_btn_sub.font.name = FONT_BODY
            p_btn_sub.font.size = Pt(8.3)
            p_btn_sub.font.color.rgb = RGBColor(50, 50, 50)

    def save(self, filepath):
        self.prs.save(filepath)
        print(f"Deck slide berhasil disimpan ke: {filepath}")

# =====================================================================
# INISIALISASI PEMBUATAN SLIDE PERTEMUAN 06 (18 SLIDES MASTERPIECE)
# =====================================================================
output_dir = r"c:\Users\anton\vibecoding\mobile2026_v2\03_slide_presentasi"
deck6 = NeoBrutalistDeckBuilder()

# ---------------------------------------------------------------------
# Slide 1: Cover
# ---------------------------------------------------------------------
deck6.add_cover(
    6,
    "State Management Reaktif: Cubit & BLoC Pattern",
    "Memisahkan Otak Logika Bisnis dari Otot Tampilan Antarmuka Menggunakan Ekosistem flutter_bloc Standar Industri Enterprise 2026.",
    "TIM DOSEN PEMROGRAMAN PIRANTI BERGERAK"
)

# ---------------------------------------------------------------------
# Slide 2: Tools & Setup flutter_bloc
# ---------------------------------------------------------------------
deck6.add_concept_with_code(
    "SETUP & ALAT BANTU",
    "Pasang Alat Tempur Cepat: flutter_bloc & Extension VS Code",
    [
        "💡 Analogi Alat: Memasang mesin pabrik modern di bengkel kita agar tidak merakit roda gigi manual dari nol.",
        "📦 Perintah Terminal: Jalankan 'flutter pub add flutter_bloc equatable' di root proyek Flutter.",
        "⚡ Ekstensi Resmi VS Code: Pasang 'Bloc' oleh Felix Angelov untuk membuat file Cubit dan State dalam 1 detik.",
        "🪄 Shortcut Ajaib: Cukup tekan Alt + Enter pada widget, pilih 'Wrap with BlocBuilder' atau 'BlocListener'."
    ],
    """// Menambahkan dependensi di pubspec.yaml:
dependencies:
  flutter:
    sdk: flutter
  flutter_bloc: ^8.1.6
  equatable: ^2.0.5

// Impor resmi di file dart:
import 'package:flutter_bloc/flutter_bloc.dart';
import 'package:equatable/equatable.dart';""",
    filename="setup_bloc_tools.dart",
    tip="Gunakan ekstensi Bloc di VS Code untuk auto-generate class Cubit dan State secara instan!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_02_tools_dan_setup_bloc.dart"
)

# -------------------------------------------------------------
# Slide 3: Ephemeral vs App State
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "THE STATE PROBLEM",
    "Dua Jenis State: Kunci Kamar Tidur vs Kunci Gerbang Kompleks",
    [
        "🔑 Analogi Kunci: Kunci kamar tidur (Ephemeral) hanya Anda yang pakai. Kunci gerbang utama (App State) dipakai bersama seluruh warga.",
        "🏠 Ephemeral State: Data lokal sementara yang hanya dibutuhkan 1 widget saja (misal: buka-tutup kartu accordion, checklist sementara).",
        "🌐 App State: Data penting yang dipakai bersama di banyak halaman (misal: item keranjang, status login, tema aplikasi).",
        "🎯 Aturan Emas: Gunakan setState() untuk Ephemeral State, dan WAJIB gunakan Cubit/BLoC untuk App State!"
    ],
    """// 1. Ephemeral State: Cukup setState biasa
bool _isCardExpanded = false;

// 2. App State: Wajib dikelola Cubit
class CartCubit extends Cubit<int> {
  CartCubit() : super(0);
  void tambahBarang() => emit(state + 1);
}""",
    filename="state_category.dart",
    tip="Jangan gunakan BLoC untuk hal sepele seperti animasi kursor; gunakan BLoC untuk data lintas halaman!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_03_ephemeral_vs_app_state.dart"
)

# -------------------------------------------------------------
# Slide 4: Keterbatasan setState() & Prop Drilling
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "LIMITASI SETSTATE",
    "Keterbatasan setState(): Bahaya Kurir Titipan Berantai",
    [
        "📦 Analogi Kurir Berantai: Kakek mau menitipkan surat ke Cucu, tapi terpaksa menitipkannya ke Ayah yang tidak peduli isi surat itu.",
        "⛓️ Prop Drilling: Mengoper variabel data dan fungsi callback melewati 5 level widget hanya agar tombol cucu bisa mengubah teks kakek.",
        "🔋 Boros Performa: Memanggil setState() di widget paling atas memaksa seluruh widget pohon di bawahnya di-rebuild tanpa alasan.",
        "🧪 Tidak Bisa Ditest: Logika perhitungan bisnis terkunci rapat di dalam kode antarmuka UI dan tidak bisa diuji unit-test murni."
    ],
    """// Kakek harus mengoper ke Ayah, Ayah mengoper ke Cucu:
LevelKakek(
  child: LevelAyah(
    counter: _counter,       // Dititipkan saja
    onIncrement: _increment, // Dititipkan saja
    child: LevelCucu(...),
  ),
);""",
    filename="prop_drilling_problem.dart",
    tip="Prop Drilling membuat kode rapuh! Sekali satu widget perantara diubah, rantai data rusak semua.",
    tag_color=COLOR_CORAL,
    full_code_file="slide_04_keterbatasan_setstate.dart"
)

# -------------------------------------------------------------
# Slide 5: Mental Model Cubit
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "ARSITEKTUR LOGIKA",
    "Mental Model Cubit: Pabrik Radio & Konsep Stream Reaktif",
    [
        "📻 Analogi Pabrik Radio: Cubit bertindak sebagai stasiun pemancar. Widget UI adalah radio penerima yang mendengarkan siaran frekuensi.",
        "🌊 Konsep Stream: Di balik layar, Cubit adalah pipa aliran data (Stream). Setiap emit() mengalirkan data baru ke widget secara instan.",
        "⚖️ Cubit vs BLoC: Cubit memakai fungsi langsung (ringkas untuk 90% aplikasi). BLoC memakai objek Event (untuk debounce pencarian).",
        "🧠 Otak vs Otot: Cubit murni kode logika Dart tanpa widget; UI murni bertugas menggambar pixel ke layar tanpa logika bisnis."
    ],
    """// Otak murni tanpa ketergantungan UI Flutter:
class CounterCubit extends Cubit<int> {
  CounterCubit() : super(0); // Nilai awal siaran

  void tambah() => emit(state + 1); // Pancarkan
  void kurang() => emit(state - 1);
}""",
    filename="mental_model_cubit.dart",
    tip="Gunakan Cubit untuk 90% fitur aplikasi! Beralihlah ke full BLoC hanya jika Anda butuh debounce pada Live Search Bar.",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_05_mental_model_cubit.dart"
)

# -------------------------------------------------------------
# Slide 6: Immutability & Equatable
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "IMMUTABLE STATE",
    "Immutability & Equatable: Stempel Cetak Ulang KTP",
    [
        "🪪 Analogi KTP: Jangan mencoret data KTP lama saat ganti alamat; cetaklah blanko KTP baru yang bersih (Immutable / Tak Berubah).",
        "🚫 Larangan Mutasi: Dilarang mengubah properti objek secara langsung (misal: 'state.poin = 100'). Flutter tidak akan mendeteksi perubahannya!",
        "✨ Method copyWith(): Cara resmi membuat salinan objek baru dengan mengganti nilai spesifik yang diinginkan.",
        "⚖️ Keajaiban Equatable: Membandingkan isi nilai objek, bukan alamat memori, sehingga mencegah rebuild duplikat yang mubazir."
    ],
    """class UserState extends Equatable {
  final String nama;
  final int poin;

  const UserState({required this.nama, required this.poin});

  UserState copyWith({String? nama, int? poin}) {
    return UserState(
      nama: nama ?? this.nama,
      poin: poin ?? this.poin,
    );
  }

  @override
  List<Object?> get props => [nama, poin];
}""",
    filename="equatable_state.dart",
    tip="Selalu extends Equatable pada class State Anda agar Flutter BLoC tahu persis kapan data benar-benar berubah!",
    tag_color=COLOR_MINT,
    full_code_file="slide_06_immutability_equatable.dart"
)

# -------------------------------------------------------------
# Slide 7: Anatomi Kelas Cubit & emit()
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "ANATOMI CUBIT",
    "Anatomi Kelas Cubit: Manajer Gudang & Pengumuman emit()",
    [
        "📢 Analogi Manajer Gudang: Saat barang baru masuk gudang, manajer mengumumkan lewat mikrofon 'emit()' ke seluruh staf kasir.",
        "🏗️ Pewarisan 'Cubit<T>': Kelas Cubit wajib mewarisi 'Cubit<NamaState>' dan memanggil 'super(initialState)' di konstruktornya.",
        "📣 Fungsi Sakti 'emit()': Satu-satunya jalan legal untuk memperbarui data state dan memberitahu UI agar menggambar ulang.",
        "🔒 Keamanan: 'emit()' berstatus protected—hanya bisa dipanggil dari dalam method internal Cubit itu sendiri."
    ],
    """class CartCubit extends Cubit<CartState> {
  CartCubit() : super(const CartState(total: 0));

  void beliBarang(int harga) {
    // Memancarkan state baru ke seluruh aplikasi:
    emit(state.copyWith(total: state.total + harga));
  }
}""",
    filename="anatomi_cubit.dart",
    tip="Dilarang emit() jika nilai sama, dan selalu cek 'if (!isClosed) emit(...)' pada operasi async agar bebas bug!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_07_anatomi_cubit_emit.dart"
)

# -------------------------------------------------------------
# Slide 8: BlocProvider & context.read()
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "INJEKSI DEPENDENSI",
    "BlocProvider & Context: Terminal Colokan Listrik Dinding",
    [
        "🔌 Analogi Colokan Listrik: Pasang stopkontak di dinding kamar (BlocProvider). Alat elektronik apa pun (Widget anak) tinggal colok.",
        "🌳 Pohon Widget: 'BlocProvider(create: (ctx) => MyCubit(), child: ...)' menempatkan instance Cubit agar hidup di ranting widget.",
        "🎯 context.read<T>(): Digunakan di dalam tombol aksi (onPressed) untuk memicu method Cubit tanpa mendengarkan perubahan.",
        "👀 context.watch<T>(): Mendengarkan perubahan data terus-menerus dan me-rebuild widget saat data baru terpancar."
    ],
    """// 1. Menyediakan Cubit ke widget anak:
BlocProvider(
  create: (context) => CounterCubit(),
  child: const CounterPage(),
);

// 2. Memanggil method Cubit saat tombol diklik:
FilledButton(
  onPressed: () => context.read<CounterCubit>().tambah(),
  child: const Icon(Icons.add),
);""",
    filename="bloc_provider_usage.dart",
    tip="Gunakan context.read di dalam callback onPressed, dan gunakan context.watch hanya di dalam method build()!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_08_bloc_provider_context.dart"
)

# -------------------------------------------------------------
# Slide 9: BlocBuilder & buildWhen
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "KONSUMSI UI 1",
    "BlocBuilder & buildWhen: Layar TV Menyala Hanya Saat Siaran Cocok",
    [
        "📺 Analogi TV Hemat Listrik: Layar TV yang hanya menyala saat acara favorit Anda mulai tayang, bukan menyala setiap detik.",
        "🎨 Peran 'BlocBuilder': Menggambar ulang widget antarmuka setiap kali Cubit memancarkan state baru.",
        "⚡ Optimasi 'buildWhen': Filter kondisi cerdas '(prev, curr) => curr % 5 == 0'. Mencegah render ulang jika syarat belum terpenuhi.",
        "🚫 Pantangan Keras: DILARANG menampilkan SnackBar, Dialog, atau Navigasi halaman di dalam BlocBuilder!"
    ],
    """BlocBuilder<StepCubit, int>(
  // Hanya rebuild jika angka kelipatan 5:
  buildWhen: (previous, current) => current % 5 == 0,
  builder: (context, step) {
    return Text(
      'Milestone Kelipatan 5: $step',
      style: const TextStyle(fontWeight: FontWeight.bold),
    );
  },
);""",
    filename="bloc_builder_filter.dart",
    tip="Bungkus hanya widget kecil yang membutuhkan data, jangan bungkus seluruh Scaffold dengan BlocBuilder!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_09_bloc_builder_buildwhen.dart"
)

# -------------------------------------------------------------
# Slide 10: BlocListener & listenWhen
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "KONSUMSI UI 2",
    "BlocListener: Bunyi Klakson Notifikasi & Pindah Halaman",
    [
        "🔔 Analogi Notifikasi WA: Anda tidak perlu menatap layar terus; saat pesan penting masuk, nada dering 'ting' berbunyi 1 kali.",
        "🎯 Penanganan 'Side Effect': Tempat resmi memanggil aksi satu kali, seperti memunculkan SnackBar, Dialog galat, atau Navigasi rute.",
        "🛡️ Bebas Duplikasi: Berbeda dengan builder yang bisa terpanggil berulang kali, listener hanya dieksekusi tepat 1 kali per transisi.",
        "🔍 Filter 'listenWhen': Hanya mendengarkan saat kondisi tertentu terjadi (misal: saat status berubah menjadi 'error' atau 'kode 200')."
    ],
    """BlocListener<FormCubit, int>(
  listenWhen: (prev, curr) => curr == 200,
  listener: (context, code) {
    ScaffoldMessenger.of(context).showSnackBar(
      const SnackBar(content: Text('Pendaftaran Berhasil!')),
    );
  },
  child: const MyFormWidget(),
);""",
    filename="bloc_listener_side_effects.dart",
    tip="Gunakan BlocListener saat ingin memindahkan halaman pengguna (context.go) setelah sukses login!",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_10_bloc_listener_side_effects.dart"
)

# -------------------------------------------------------------
# Slide 11: BlocConsumer
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "KONSUMSI UI 3",
    "BlocConsumer: Duo Komentator & Papan Skor Pertandingan",
    [
        "⚽ Analogi Stadion Bola: Papan skor mencatat gol (builder), sementara komentator mengumumkan gol lewat peluit (listener).",
        "🤝 Dua Fungsi Bersatu: Menggabungkan 'BlocBuilder' dan 'BlocListener' ke dalam satu widget ringkas tanpa tumpukan bersarang.",
        "📋 Parameter 'listener': Untuk menampilkan pesan feedback SnackBar atau navigasi.",
        "🖼️ Parameter 'builder': Untuk menampilkan tampilan visual tombol, teks skor, atau ikon status."
    ],
    """BlocConsumer<LoginCubit, bool>(
  listener: (context, isLoggedIn) {
    if (isLoggedIn) {
      ScaffoldMessenger.of(context).showSnackBar(
        const SnackBar(content: Text('Selamat Datang!')),
      );
    }
  },
  builder: (context, isLoggedIn) {
    return FilledButton(
      onPressed: () => context.read<LoginCubit>().toggle(),
      child: Text(isLoggedIn ? 'Logout' : 'Login'),
    );
  },
);""",
    filename="bloc_consumer_duo.dart",
    tip="Gunakan BlocConsumer pada tombol autentikasi login atau checkout keranjang belanja!",
    tag_color=COLOR_MINT,
    full_code_file="slide_11_bloc_consumer_duo.dart"
)

# -------------------------------------------------------------
# Slide 12: BlocSelector
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "OPTIMASI REBUILD",
    "BlocSelector: Lensa Kamera Zoom ke Objek Tertentu",
    [
        "🔬 Analogi Lensa Zoom: Memotret burung di dahan pohon tanpa peduli daun di sekitarnya bergoyang tertiup angin.",
        "🎯 Ekstraksi Properti Spesifik: Jika State Anda memiliki 10 properti, 'BlocSelector' hanya memantau 1 properti saja.",
        "⚡ Efisiensi Tingkat Dewa: Widget nama pengguna tidak akan pernah di-rebuild meskipun umur atau saldo pengguna berubah ribuan kali.",
        "🏭 Skala Enterprise: Kunci utama aplikasi Flutter skala besar tetap berjalan mulus 60-120 FPS tanpa patah-patah (*stutter*)."
    ],
    """BlocSelector<ProfilCubit, ProfilData, String>(
  selector: (state) => state.nama, // Hanya pantau 'nama'
  builder: (context, nama) {
    return Text('Nama Pengguna: $nama');
    // Tidak akan di-rebuild saat properti 'umur' berubah!
  },
);""",
    filename="bloc_selector_opt.dart",
    tip="BlocSelector adalah senjata rahasia saat aplikasi memiliki formulir raksasa dengan puluhan field input!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_12_bloc_selector_opt.dart"
)

# -------------------------------------------------------------
# Slide 13: Arsitektur 4 Status UI
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "POLA STATUS UI",
    "Arsitektur 4 Status UI: Empat Lampu Indikator Mesin Cuci",
    [
        "🧺 Analogi Mesin Cuci: Siap (Initial), Sedang Mencuci (Loading), Cucian Selesai (Success), atau Air Mampet (Error).",
        "🏛️ Pola Baku Industri: Seluruh interaksi data dan API wajib dimodelkan ke dalam 4 state terpisah yang jelas (*sealed class*).",
        "⏳ DataInitial: Kondisi awal saat layar baru dibuka sebelum aksi apa pun.",
        "🔄 DataLoading & DataSuccess / DataError: Menjamin pengguna tidak pernah bingung melihat layar kosong (*Blank Screen*)."
    ],
    """sealed class DataState extends Equatable {}

class DataInitial extends DataState { ... }
class DataLoading extends DataState { ... }
class DataSuccess extends DataState { final List items; ... }
class DataError extends DataState { final String pesan; ... }

// Di UI, tangani dengan pola switch ekspresif:
return switch (state) {
  DataInitial() => const Text('Tekan Tombol'),
  DataLoading() => const CircularProgressIndicator.adaptive(),
  DataSuccess(:final items) => ListView(...),
  DataError(:final pesan) => Text('Galat: $pesan'),
};""",
    filename="four_ui_states.dart",
    tip="Gunakan pola pattern matching switch Dart modern untuk memastikan seluruh status UI tertangani!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_13_arsitektur_4_status_ui.dart"
)

# -------------------------------------------------------------
# Slide 14: MultiBlocProvider
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "MULTI-PROVIDER",
    "MultiBlocProvider: Panel Saklar Listrik Rumah Terpusat",
    [
        "⚡ Analogi Panel MCB: Di meteran listrik rumah tangga, ada saklar untuk kulkas, AC, dan lampu di dalam 1 kotak rapi.",
        "🏢 Kebutuhan Nyata: Aplikasi riil membutuhkan banyak Cubit sekaligus (AuthCubit, CartCubit, ThemeCubit, NotifCubit).",
        "🚫 Mencegah Piramida Neraka: Menghindari susunan bersarang 'BlocProvider di dalam BlocProvider di dalam BlocProvider'.",
        "🌐 Pasang di Root: Tempatkan 'MultiBlocProvider' di atas MaterialApp agar semua halaman memiliki akses data yang sama."
    ],
    """MultiBlocProvider(
  providers: [
    BlocProvider<AuthCubit>(create: (ctx) => AuthCubit()),
    BlocProvider<CartCubit>(create: (ctx) => CartCubit()),
    BlocProvider<ThemeCubit>(create: (ctx) => ThemeCubit()),
  ],
  child: const MainApp(),
);""",
    filename="multi_bloc_provider.dart",
    tip="Daftarkan seluruh Cubit global di MultiBlocProvider pada file main.dart aplikasi Anda!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_14_multi_bloc_provider.dart"
)

# -------------------------------------------------------------
# Slide 15: Pemantauan Global dengan BlocObserver
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "GLOBAL MONITORING",
    "BlocObserver: CCTV Ruang Satpam Pemantau Seluruh Gerakan",
    [
        "📹 Analogi CCTV Satpam: Kamera pemantau yang mencatat siapa pun yang keluar-masuk gedung tanpa mengganggu aktivitas staf.",
        "🕵️ Melacak Perubahan Global: 'BlocObserver' menangkap setiap kali ada Cubit yang memanggil 'emit()' di seluruh aplikasi.",
        "🖥️ Logging Otomatis di Terminal: Mencetak 'CartCubit: state lama -> state baru' secara otomatis tanpa menulis 'print()' manual.",
        "🚨 Deteksi Galat Terpusat: Method 'onError()' menangkap error logika unhandled sebelum aplikasi sempat mengalami crash."
    ],
    """class AppBlocObserver extends BlocObserver {
  @override
  void onChange(BlocBase bloc, Change change) {
    super.onChange(bloc, change);
    debugPrint('[CCTV] ${bloc.runtimeType}: '
        '${change.currentState} -> ${change.nextState}');
  }
}

void main() {
  // Pasang CCTV sebelum aplikasi berjalan:
  Bloc.observer = AppBlocObserver();
  runApp(const MyApp());
}""",
    filename="bloc_observer_cctv.dart",
    tip="BlocObserver menghemat waktu debugging Anda hingga 80% saat mencari sumber kesalahan logika data!",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_15_bloc_observer_cctv.dart"
)

# -------------------------------------------------------------
# Slide 16: Refactoring setState ke Cubit
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "STUDI KASUS REFACTOR",
    "Refactoring Nyata: Operasi Bedah Memisahkan Otak dari Otot",
    [
        "🏥 Analogi Operasi Medis: Mengambil tumor logika yang menempel di organ UI dan memindahkannya ke wadah steril tersendiri.",
        "🗑️ Bersihkan StatefulWidget: Ubah widget menjadi 'StatelessWidget' sederhana yang murni bertugas menggambar antarmuka.",
        "📦 Bungkus Data ke Cubit: Pindahkan variabel 'Set<String> wishlist' dan fungsi 'toggle()' ke dalam file 'WishlistCubit'.",
        "✨ Hasil Akhir: UI menjadi sangat bersih, reaktif, mudah dibaca rekan tim, dan siap dihubungkan ke database online!"
    ],
    """// SEBELUM (setState campur aduk di UI):
void _toggle(String id) {
  setState(() => _items.contains(id) ? _items.remove(id) : _items.add(id));
}

// SESUDAH (Cubit bersih & modular):
class WishlistCubit extends Cubit<Set<String>> {
  WishlistCubit() : super({});
  void toggle(String id) {
    final updated = Set<String>.from(state);
    updated.contains(id) ? updated.remove(id) : updated.add(id);
    emit(updated);
  }
}""",
    filename="refactoring_setstate_cubit.dart",
    tip="Refactoring dari setState ke Cubit adalah ujian keterampilan nomor 1 pada sesi Live Code Defense UTS!",
    tag_color=COLOR_MINT,
    full_code_file="slide_16_refactoring_setstate_cubit.dart"
)

# -------------------------------------------------------------
# Slide 17: Checklist 5 Aturan Emas State Management 2026
# -------------------------------------------------------------
deck6.add_concept_with_code(
    "CHECKLIST PRODUKSI",
    "5 Aturan Emas State Management (Sabuk Pengaman Developer)",
    [
        "1. State Wajib Immutable: Selalu extends Equatable dan gunakan copyWith() untuk membuat data baru.",
        "2. Dilarang Memanggil Side Effect di BlocBuilder: Gunakan BlocListener khusus untuk SnackBar & Navigasi.",
        "3. Gunakan context.read() di Event Callback: Hanya gunakan context.watch() saat membaca nilai untuk tampilan UI.",
        "4. Pasang BlocObserver di main.dart: Jangan pernah merilis aplikasi tanpa CCTV pencatat transisi state.",
        "5. Pisahkan File State & Cubit: Satu fitur wajib memiliki folder /cubit/ terisolasi (Feature-First)."
    ],
    """// Checklist Uji Praktikum State Management Standar 2026:
// [✓] Seluruh App State dikelola Cubit (Zero setState lintas widget)
// [✓] BlocBuilder tidak memicu SnackBar/Dialog galat
// [✓] Equatable mencegah render ulang saat data tidak berubah
// [✓] Log transisi tercetak rapi di terminal via BlocObserver""",
    filename="checklist_state_management.dart",
    tip="Jadikan kelima checklist ini sebagai acuan penilaian proyek UTS kelompok Anda agar meraih nilai A!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_17_checklist_state_management.dart"
)

# -------------------------------------------------------------
# Slide 18: Lab Quest 06 - Cart & Wishlist dengan Cubit
# -------------------------------------------------------------
deck6.add_lab_quest(
    6,
    "Aplikasi Mini E-Commerce Multi-Bloc (Cart & SnackBar Reaktif)",
    60,
    [
        "Bangun aplikasi Toko Piranti Bergerak dengan State Management Cubit terpisah dari antarmuka UI.",
        "Buat kelas 'CartCubit' dan 'CartItemState' yang mewarisi Equatable untuk mencatat daftar produk yang dibeli.",
        "Tampilkan Badge jumlah item keranjang di AppBar yang otomatis bertambah saat tombol 'Beli' diklik.",
        "Terapkan 'BlocListener' untuk memunculkan SnackBar hijau 'Berhasil menambah barang' setiap kali item baru masuk keranjang.",
        "Sediakan BottomBar yang menghitung total harga belanjaan secara otomatis dan tombol 'Checkout / Reset'.",
        "Target Pengujian: Jalankan di browser Chrome ('flutter run -d chrome'), buktikan bahwa penambahan item berjalan instan tanpa lag!"
    ],
    "Keranjang belanja bertambah reaktif, badge AppBar otomatis terupdate, SnackBar muncul via BlocListener, dan total harga akurat.",
    full_code_file="slide_18_lab_quest_cart_cubit.dart"
)

# Simpan berkas presentasi
output_pptx = os.path.join(output_dir, "PERTEMUAN_06_State_Management_Cubit_dan_BLoC.pptx")
deck6.save(output_pptx)
