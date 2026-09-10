"""
Generator Slide Pertemuan 04 (Edisi Masterpiece 19 Slide Lengkap & Super Ramah Awam)
Mata Kuliah: Pemrograman Piranti Bergerak (Flutter & Dart 2026)
Topik: Form Handling, Input Validation, & Navigasi Deklaratif GoRouter
Standar Desain: Neo-Brutalism Murni (designprompts.dev/neo-brutalism)
- ZERO BORDER RADIUS: Sudut 90 derajat siku murni tanpa kelengkungan (MSO_SHAPE.RECTANGLE).
- BOLD SOLID BLACK BORDERS: Garis hitam tegas 2.5 Pt (#000000).
- HARD OFFSET BLOCK SHADOWS: Bayangan blok hitam solid +0.08 inch di belakang kartu.
- PALET NEO-BRUTALISM: Warm Cream (#FAF8F5), Canary Yellow (#FFE600), Electric Cyan (#38BDF8), Mint (#4ADE80), Coral (#FB7185), Lavender (#C084FC).
- FITUR UTAMA: 100% SLIDE PEMBAHASAN DISERTAI KODE NYATA FLUTTER, ANALOGI MEMBUMI, DAN TIPS AWAM.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- PALET WARNA NEO-BRUTALISM ---
COLOR_BG = RGBColor(250, 248, 245)        # #FAF8F5 (Warm Cream Canvas)
COLOR_BLACK = RGBColor(0, 0, 0)           # #000000 (Pure Black Border & Hard Shadows)
COLOR_WHITE = RGBColor(255, 255, 255)     # #FFFFFF (Pure White Card)
COLOR_YELLOW = RGBColor(255, 230, 0)      # #FFE600 (Canary Yellow)
COLOR_CYAN = RGBColor(56, 189, 248)       # #38BDF8 (Electric Cyan)
COLOR_MINT = RGBColor(74, 222, 128)       # #4ADE80 (Neo Mint / Green)
COLOR_CORAL = RGBColor(251, 113, 133)     # #FB7185 (Neo Coral / Pink)
COLOR_PURPLE = RGBColor(192, 132, 252)    # #C084FC (Lavender Purple)
COLOR_CODE_BG = RGBColor(18, 18, 18)      # #121212 (Deep Black Code Box)
COLOR_CODE_TEXT = RGBColor(245, 245, 245) # #F5F5F5 (Crisp Monospace)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

SHADOW_OFFSET = Inches(0.08)

GITHUB_BASE_URL = "https://github.com/antonprafanto/mobile2026/blob/main/07_contoh_kode_program/pertemuan_04_form_dan_gorouter"

class NeoBrutalistDeckBuilder:
    def __init__(self, course_name="IF3205 • PEMROGRAMAN PIRANTI BERGERAK"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.course_name = course_name

    def _set_canvas_bg(self, slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def _add_neobrutal_card(self, slide, x, y, w, h, fill_color=COLOR_WHITE, has_shadow=True, border_width=Pt(2.5)):
        if has_shadow:
            shadow = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, x + SHADOW_OFFSET, y + SHADOW_OFFSET, w, h
            )
            shadow.fill.solid()
            shadow.fill.fore_color.rgb = COLOR_BLACK
            shadow.line.fill.background()

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, y, w, h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = COLOR_BLACK
        card.line.width = border_width
        return card

    def _add_header(self, slide, tag, title, tag_color=COLOR_YELLOW):
        tag_w = Inches(len(tag) * 0.125 + 0.6)
        self._add_neobrutal_card(slide, Inches(0.9), Inches(0.4), tag_w, Inches(0.4), fill_color=tag_color, has_shadow=True)
        tb_tag = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), tag_w, Inches(0.4))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.name = FONT_HEADING
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_BLACK

        tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.88), Inches(11.5), Inches(0.8))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_BLACK

    # 1. Slide Cover
    def add_cover(self, meeting_num, title, subtitle, presenter_info):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)

        self._add_neobrutal_card(slide, Inches(0.9), Inches(0.9), Inches(11.533), Inches(5.7), fill_color=COLOR_WHITE)

        tb = slide.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.5), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"✦ {self.course_name.upper()} • PERTEMUAN {meeting_num:02d} ✦"
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_BLACK
        p0.space_after = Pt(18)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(30)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BLACK
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = FONT_BODY
        p2.font.size = Pt(15)
        p2.font.color.rgb = RGBColor(60, 60, 60)
        p2.space_after = Pt(38)

        p3 = tf.add_paragraph()
        p3.text = f"■  {presenter_info}  ■  OUTCOME-BASED EDUCATION (OBE) • SUB-CPMK 4"
        p3.font.name = FONT_HEADING
        p3.font.size = Pt(12)
        p3.font.bold = True
        p3.font.color.rgb = COLOR_BLACK

    # 2. Slide Split Screen: Konsep & Tips di Kiri + Kode Sederhana di Kanan
    def add_concept_with_code(self, tag, title, bullets, code_snippet, filename="form_handling.dart", tip=None, tag_color=COLOR_YELLOW, full_code_file=None):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, tag, title, tag_color=tag_color)

        # Left Explanations Card (4.9 Inches)
        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(4.9), Inches(5.1), fill_color=COLOR_WHITE)

        tb_l = slide.shapes.add_textbox(Inches(1.15), Inches(1.95), Inches(4.4), Inches(4.8))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True

        p_lt = tf_l.paragraphs[0]
        p_lt.text = "KONSEP & ALUR KERJA"
        p_lt.font.name = FONT_HEADING
        p_lt.font.size = Pt(11.5)
        p_lt.font.bold = True
        p_lt.font.color.rgb = COLOR_BLACK
        p_lt.space_after = Pt(8)

        for b in bullets:
            p = tf_l.add_paragraph()
            p.text = f"■  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.6)
            p.font.color.rgb = RGBColor(20, 20, 20)
            p.space_after = Pt(5)

        if tip:
            p_tip = tf_l.add_paragraph()
            p_tip.text = f"💡 Tips Praktis: {tip}"
            p_tip.font.name = FONT_HEADING
            p_tip.font.size = Pt(10.2)
            p_tip.font.bold = True
            p_tip.font.color.rgb = COLOR_BLACK

        # Right Sharp Black Code Box (6.433 Inches)
        self._add_neobrutal_card(slide, Inches(6.0), Inches(1.8), Inches(6.433), Inches(5.1), fill_color=COLOR_CODE_BG)

        # Header Code Strip
        c_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), Inches(1.8), Inches(6.433), Inches(0.48))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = tag_color
        c_bar.line.color.rgb = COLOR_BLACK
        c_bar.line.width = Pt(2.5)

        tb_cb = slide.shapes.add_textbox(Inches(6.2), Inches(1.86), Inches(6.0), Inches(0.38))
        p_cb = tb_cb.text_frame.paragraphs[0]
        p_cb.text = f"{filename}  |  KODE IMPLEMENTASI"
        p_cb.font.name = FONT_CODE
        p_cb.font.size = Pt(10.5)
        p_cb.font.bold = True
        p_cb.font.color.rgb = COLOR_BLACK

        code_box_h = Inches(3.68) if full_code_file else Inches(4.35)
        tb_r = slide.shapes.add_textbox(Inches(6.2), Inches(2.32), Inches(6.0), code_box_h)
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_code = tf_r.paragraphs[0]
        p_code.text = code_snippet
        p_code.font.name = FONT_CODE
        p_code.font.size = Pt(9.1)
        p_code.font.color.rgb = COLOR_CODE_TEXT

        # Neo-Brutalist Runnable Code Hyperlink Button at Bottom of Right Box
        if full_code_file:
            btn_x = Inches(6.15)
            btn_y = Inches(6.16)
            btn_w = Inches(6.133)
            btn_h = Inches(0.60)

            # Button Shadow (Neo-Brutalism offset)
            btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.04), btn_y + Inches(0.04), btn_w, btn_h)
            btn_shadow.fill.solid()
            btn_shadow.fill.fore_color.rgb = COLOR_BLACK
            btn_shadow.line.fill.background()

            # Button Face
            btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
            btn.fill.solid()
            btn.fill.fore_color.rgb = COLOR_YELLOW
            btn.line.color.rgb = COLOR_BLACK
            btn.line.width = Pt(2.0)

            tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.08), btn_y + Inches(0.03), btn_w - Inches(0.16), btn_h - Inches(0.06))
            tf_btn = tb_btn.text_frame
            tf_btn.word_wrap = True

            p_btn = tf_btn.paragraphs[0]
            run_btn = p_btn.add_run()
            run_btn.text = f"▶ BUKA KODE LENGKAP DI GITHUB ({full_code_file})"
            run_btn.font.name = FONT_HEADING
            run_btn.font.size = Pt(9.6)
            run_btn.font.bold = True
            run_btn.font.color.rgb = COLOR_BLACK
            run_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

            p_btn_sub = tf_btn.add_paragraph()
            p_btn_sub.text = f"💡 Buka di browser / salin source code: github.com/antonprafanto/mobile2026"
            p_btn_sub.font.name = FONT_BODY
            p_btn_sub.font.size = Pt(7.6)
            p_btn_sub.font.color.rgb = RGBColor(60, 60, 60)

    # 3. Slide Lab Quest
    def add_lab_quest(self, meeting_num, title, time_minutes, goals, success_criteria, full_code_file="slide_19_lab_quest_portal_ktm.dart"):
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._set_canvas_bg(slide)
        self._add_header(slide, "LAB QUEST MANDIRI", f"Pertemuan {meeting_num:02d}: {title}", tag_color=COLOR_YELLOW)

        self._add_neobrutal_card(slide, Inches(0.9), Inches(1.8), Inches(11.533), Inches(5.1), fill_color=COLOR_WHITE)

        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(11.533), Inches(0.52))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_YELLOW
        banner.line.color.rgb = COLOR_BLACK
        banner.line.width = Pt(2.5)

        tb_b = slide.shapes.add_textbox(Inches(1.2), Inches(1.85), Inches(10.9), Inches(0.45))
        p_b = tb_b.text_frame.paragraphs[0]
        p_b.text = f"⏱️ ALOKASI: {time_minutes} MENIT  ■  TARGET UJI: BROWSER CHROME ATAU SMARTPHONE FISIK"
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(11.5)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_BLACK

        tb = slide.shapes.add_textbox(Inches(1.3), Inches(2.40), Inches(10.7), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_g = tf.paragraphs[0]
        p_g.text = "LANGKAH KERJA & SPESIFIKASI TUGAS:"
        p_g.font.name = FONT_HEADING
        p_g.font.size = Pt(12)
        p_g.font.bold = True
        p_g.font.color.rgb = COLOR_BLACK
        p_g.space_after = Pt(6)

        for g in goals:
            p = tf.add_paragraph()
            p.text = f"✔  {g}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.4)
            p.font.color.rgb = RGBColor(30, 30, 30)
            p.space_after = Pt(4)

        p_sc = tf.add_paragraph()
        p_sc.text = f"🎯 TARGET SELESAI: {success_criteria}"
        p_sc.font.name = FONT_HEADING
        p_sc.font.size = Pt(10.8)
        p_sc.font.bold = True
        p_sc.font.color.rgb = RGBColor(10, 100, 40)

        # Button for Lab Quest Source Code
        if full_code_file:
            btn_x = Inches(1.3)
            btn_y = Inches(6.12)
            btn_w = Inches(10.7)
            btn_h = Inches(0.62)

            btn_shadow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x + Inches(0.04), btn_y + Inches(0.04), btn_w, btn_h)
            btn_shadow.fill.solid()
            btn_shadow.fill.fore_color.rgb = COLOR_BLACK
            btn_shadow.line.fill.background()

            btn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, btn_x, btn_y, btn_w, btn_h)
            btn.fill.solid()
            btn.fill.fore_color.rgb = COLOR_YELLOW
            btn.line.color.rgb = COLOR_BLACK
            btn.line.width = Pt(2.0)

            tb_btn = slide.shapes.add_textbox(btn_x + Inches(0.12), btn_y + Inches(0.03), btn_w - Inches(0.24), btn_h - Inches(0.06))
            tf_btn = tb_btn.text_frame
            tf_btn.word_wrap = True

            p_btn = tf_btn.paragraphs[0]
            run_btn = p_btn.add_run()
            run_btn.text = f"▶ BUKA KODE LENGKAP LAB QUEST DI GITHUB ({full_code_file})"
            run_btn.font.name = FONT_HEADING
            run_btn.font.size = Pt(11)
            run_btn.font.bold = True
            run_btn.font.color.rgb = COLOR_BLACK
            run_btn.hyperlink.address = f"{GITHUB_BASE_URL}/{full_code_file}"

            p_btn_sub = tf_btn.add_paragraph()
            p_btn_sub.text = "💡 Buka di browser: Kode lengkap Form Registrasi, Validasi Regex, Unfocus Keyboard, & KTM Digital di GitHub"
            p_btn_sub.font.name = FONT_BODY
            p_btn_sub.font.size = Pt(8.3)
            p_btn_sub.font.color.rgb = RGBColor(50, 50, 50)

    def save(self, filepath):
        self.prs.save(filepath)
        print(f"Deck slide berhasil disimpan ke: {filepath}")

# =====================================================================
# INISIALISASI PEMBUATAN SLIDE PERTEMUAN 04 (19 SLIDES MASTERPIECE)
# =====================================================================
output_dir = r"c:\Users\anton\vibecoding\mobile2026_v2\03_slide_presentasi"
deck4 = NeoBrutalistDeckBuilder()

# ---------------------------------------------------------------------
# Slide 1: Cover Pertemuan 04
# ---------------------------------------------------------------------
deck4.add_cover(
    4,
    "Form Handling, Input Validation, & Navigasi Deklaratif GoRouter",
    "Membangun Formulir Interaktif Bebas Kebocoran Memori, Validasi Regex Otomatis, dan Navigasi URL Modern Standar Industri 2026.",
    "TIM DOSEN PEMROGRAMAN PIRANTI BERGERAK"
)

# ---------------------------------------------------------------------
# Slide 2: Panduan Alat Pengujian & Instalasi Cepat (Ramah Awam)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "SETUP & TOOLS",
    "Persiapan Uji Cepat Hari Ini: Terminal VS Code & Chrome",
    [
        "Sebelum mulai ngoding: Buka terminal terintegrasi di VS Code dengan shortcut 'Ctrl + ~' (Control + Backtick).",
        "Pasang library navigasi resmi: Jalankan perintah 'flutter pub add go_router'. Pustaka terpasang otomatis tanpa risiko salah indentasi di pubspec.yaml!",
        "Target Pengujian Cepat: Ketik 'flutter run -d chrome' untuk menjalankan di peramban web Chrome (hemat RAM lab & instan 5 detik).",
        "Pilihan Smartphone: Hubungkan kabel USB ponsel fisik dan buka software 'scrcpy' untuk sensasi papan ketik HP asli."
    ],
    """# 1. Buka terminal di VS Code (tekan Ctrl + ~)

# 2. Pasang library GoRouter versi resmi terbaru:
flutter pub add go_router

# 3. Jalankan aplikasi di browser Chrome (Paling ringan lab PC):
flutter run -d chrome

# Hasil otomatis di berkas 'pubspec.yaml':
# dependencies:
#   flutter:
#     sdk: flutter
#   go_router: ^14.8.0  # Terpasang otomatis rapi!""",
    filename="terminal_setup.sh",
    tip="Gunakan 'flutter pub add' di terminal agar berkas pubspec.yaml Anda tidak rusak akibat salah spasi!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_02_persiapan_tools_dan_gorouter.dart"
)

# ---------------------------------------------------------------------
# Slide 3: Dua Jenis Input: TextField vs TextFormField
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "KOMPONEN INPUT",
    "Dua Jenis Input Teks: TextField vs TextFormField",
    [
        "Kapan memakai TextField biasa? Sangat cocok untuk input mandiri tanpa validasi (misal: kolom pencarian 'Search Bar' di AppBar).",
        "Kapan memakai TextFormField? Wajib digunakan di formulir pendaftaran/login yang butuh pengecekan data sebelum dikirim.",
        "Keunggulan TextFormField: Jika isian salah, teks peringatan merah (error text) otomatis muncul tepat di bawah kotak input tanpa perlu setState manual.",
        "Terintegrasi dengan Form: Mendukung fungsi pengecekan 'validator' dan dapat diperiksa serentak via '_formKey.currentState!.validate()'.",
    ],
    """// 1. TextField Sederhana (Untuk Kolom Pencarian):
TextField(
  decoration: const InputDecoration(
    hintText: 'Cari nama mahasiswa...',
    prefixIcon: Icon(Icons.search),
  ),
  onChanged: (text) => print('Mencari: $text'),
);

// 2. TextFormField Cerdas (Untuk Formulir Validasi):
TextFormField(
  decoration: const InputDecoration(labelText: 'Email Kampus'),
  validator: (val) {
    if (val == null || val.isEmpty) return 'Email wajib diisi!';
    return null; // Valid!
  },
);""",
    filename="text_field_comparison.dart",
    tip="Jika ada tombol Submit/Kirim data, SELALU gunakan TextFormField!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_03_textfield_vs_textformfield.dart"
)

# ---------------------------------------------------------------------
# Slide 4: Mengendalikan Input: TextEditingController (Anti-Bocor Memori)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "KENDALI INPUT",
    "Membaca & Membersihkan Input: TextEditingController",
    [
        "TextEditingController bertindak sebagai remote pengendali kolom input: membaca isi teks (_ctrl.text), mengubah teks, atau mengosongkan kolom (_ctrl.clear()).",
        "Wajib diletakkan di dalam StatefulWidget pada variabel kelas State.",
        "⚠️ Peringatan Kritis Mahasiswa Awam: Controller WAJIB dimatikan menggunakan method '_ctrl.dispose()' di dalam blok dispose().",
        "Jika lupa di-dispose, objek controller akan terus tertinggal di RAM dan memicu kebocoran memori (Memory Leak)!"
    ],
    """class InputDemoState extends State<InputDemo> {
  // 1. Deklarasi Controller di awal kelas State:
  final _namaCtrl = TextEditingController();

  @override
  void dispose() {
    // 2. WAJIB dibersihkan saat halaman ditutup:
    _namaCtrl.dispose();
    super.dispose();
  }

  void _bacaInput() {
    // 3. Baca teks yang diketik pengguna:
    print('Nama yang diketik: ${_namaCtrl.text}');
  }
}""",
    filename="controller_lifecycle.dart",
    tip="Hafalkan aturan emas: Setiap TextEditingController yang Anda buat WAJIB di-dispose()!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_04_controller_lifecycle.dart"
)

# ---------------------------------------------------------------------
# Slide 5: Desain UX Input: InputDecoration, Hint, & Tipe Keyboard HP
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "DESAIN PENGALAMAN INPUT",
    "Mempercantik Input: InputDecoration & Tipe Keyboard HP",
    [
        "labelText: Label teks mengambang yang naik ke atas saat pengguna mulai mengetik.",
        "hintText: Contoh isian samar penunjuk format (misal: 'nama@mhs.kampus.ac.id').",
        "prefixIcon: Ikon pemanis di awal kolom (amplop untuk email, gembok untuk sandi).",
        "border: OutlineInputBorder() untuk memberi batas kotak tegas bergaya modern.",
        "keyboardType: Menyesuaikan papan ketik HP pengguna (angka murni, email dengan tombol '@', atau telepon)."
    ],
    """TextFormField(
  controller: _emailCtrl,
  keyboardType: TextInputType.emailAddress, // Papan ketik ber-simbol @
  decoration: const InputDecoration(
    labelText: 'Email Mahasiswa',
    hintText: 'nama@mhs.kampus.ac.id',
    prefixIcon: Icon(Icons.email_outlined),
    border: OutlineInputBorder(), // Border kotak tegas
  ),
);""",
    filename="input_decoration_ux.dart",
    tip="Mengatur keyboardType secara tepat menghindarkan pengguna HP dari kerumitan mencari simbol '@'!",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_05_input_decoration_ux.dart"
)

# ---------------------------------------------------------------------
# Slide 6: Interaktivitas Sandi: Fitur Intip Password Nyata (_isObscure)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "FITUR KEAMANAN UX",
    "Toggle Sandi: Membuat Tombol Intip Password yang Berfungsi",
    [
        "Masalah Pemula: Tahu obscureText: true, tapi bingung cara membuat ikon mata bisa diklik untuk melihat sandi.",
        "Solusi Mudah: Buat variabel boolean '_isObscure = true' di kelas State.",
        "Pasang IconButton pada 'suffixIcon' dengan ikon kondisional: jika true tampilkan mata terbuka, jika false tampilkan mata dicoret.",
        "Panggil 'setState(() => _isObscure = !_isObscure)' saat ikon ditekan!"
    ],
    """class _PasswordInputState extends State<PasswordInput> {
  bool _isObscure = true; // Status awal tersembunyi

  @override
  Widget build(BuildContext context) {
    return TextFormField(
      obscureText: _isObscure, // Dikendalikan boolean
      decoration: InputDecoration(
        labelText: 'Kata Sandi',
        prefixIcon: const Icon(Icons.lock_outline),
        suffixIcon: IconButton(
          icon: Icon(_isObscure ? Icons.visibility : Icons.visibility_off),
          onPressed: () {
            setState(() => _isObscure = !_isObscure); // Toggle bolak-balik
          },
        ),
      ),
    );
  }
}""",
    filename="password_toggle_demo.dart",
    tip="Fitur toggle intip sandi adalah standar wajib pada semua formulir login komersial modern!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_06_password_toggle_demo.dart"
)

# ---------------------------------------------------------------------
# Slide 7: Menjinakkan Keyboard HP: Unfocus & Pindah Kolom Otomatis
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "KENDALI KEYBOARD HP",
    "Menutup Keyboard HP & Pindah Kolom Otomatis",
    [
        "Keluhan Populer Mahasiswa: 'Pak, setelah klik Submit, kenapa keyboard virtual HP tetap muncul dan menutupi tombol?'.",
        "Solusi Menutup Keyboard: Panggil 'FocusScope.of(context).unfocus()' tepat sebelum aksi submit.",
        "textInputAction: TextInputAction.next: Mengubah tombol pojok kanan keyboard HP menjadi panah 'Next' untuk langsung pindah ke kolom berikutnya.",
        "Mencegah frustrasi pengguna smartphone berlayar kecil!"
    ],
    """// 1. Pindah otomatis ke kolom berikutnya saat tekan Enter HP:
TextFormField(
  textInputAction: TextInputAction.next, // Ikon 'Next'
  decoration: const InputDecoration(labelText: 'Nama Lengkap'),
);

// 2. Tutup keyboard saat tombol Submit ditekan:
void _onSubmit() {
  // Sembunyikan keyboard virtual seketika:
  FocusScope.of(context).unfocus();

  if (_formKey.currentState!.validate()) {
    // Lanjutkan proses pendaftaran...
  }
}""",
    filename="keyboard_focus_control.dart",
    tip="Selalu panggil 'FocusScope.of(context).unfocus()' agar SnackBar dan tombol tidak terhalang keyboard!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_07_keyboard_focus_control.dart"
)

# ---------------------------------------------------------------------
# Slide 8: Input Pilihan Form: DropdownButtonFormField
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "INPUT PILIHAN",
    "Pilihan Terintegrasi Form: DropdownButtonFormField",
    [
        "Formulir pendaftaran tidak hanya berisi teks, tapi juga pilihan pasti (misal: Program Studi, Jenis Kelamin).",
        "DropdownButtonFormField: Widget dropdown bawaan Flutter yang sudah terintegrasi langsung dengan ekosistem validasi Form!",
        "Mendukung properti 'validator' layaknya TextFormField biasa.",
        "Menghindarkan pengguna dari salah ketik nama jurusan/program studi."
    ],
    """String? _selectedProdi;
final List<String> _daftarProdi = ['Informatika', 'Sistem Informasi', 'Teknik Komputer'];

DropdownButtonFormField<String>(
  value: _selectedProdi,
  decoration: const InputDecoration(labelText: 'Program Studi'),
  items: _daftarProdi.map((p) => DropdownMenuItem(value: p, child: Text(p))).toList(),
  onChanged: (val) => setState(() => _selectedProdi = val),
  validator: (val) {
    if (val == null || val.isEmpty) return 'Pilih Program Studi Anda!';
    return null; // Valid!
  },
);""",
    filename="dropdown_form_field.dart",
    tip="DropdownButtonFormField otomatis menampilkan garis merah galat jika mahasiswa belum memilih opsi!",
    tag_color=COLOR_MINT,
    full_code_file="slide_08_dropdown_form_field.dart"
)

# ---------------------------------------------------------------------
# Slide 9: Pondasi Keamanan: Form Widget & GlobalKey<FormState>
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "VALIDASI TERPUSAT",
    "Arsitektur Form: Analogi Map Berkas & Stempel Legalisir",
    [
        "Analogi: Bayangkan widget Form adalah Map Berkas yang menampung seluruh formulir kertas.",
        "GlobalKey<FormState> adalah Stempel Legalisir di Meja Registrasi.",
        "Hanya dengan 1 perintah: '_formKey.currentState!.validate()', Flutter secara otomatis memeriksa SELURUH kolom di dalam map berkas sekaligus!",
        "Jika ada satu saja kolom yang salah atau kosong, proses submit ditolak dan pesan teks merah otomatis muncul di bawah kolom tersebut."
    ],
    """class RegistrasiPage extends StatefulWidget { ... }
class _RegistrasiPageState extends State<RegistrasiPage> {
  // Stempel Kunci Form Terpusat:
  final _formKey = GlobalKey<FormState>();

  void _kirimFormulir() {
    // Periksa keabsahan seluruh kolom sekaligus:
    if (_formKey.currentState!.validate()) {
      print('Formulir LULUS validasi! Siap kirim ke server.');
    }
  }

  @override
  Widget build(BuildContext context) {
    return Form(
      key: _formKey, // Pasang kunci stempel di sini
      child: Column(children: [ ... ]),
    );
  }
}""",
    filename="form_architecture.dart",
    tip="Selalu simpan GlobalKey<FormState> sebagai variabel final di kelas State!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_09_form_architecture.dart"
)

# ---------------------------------------------------------------------
# Slide 10: Logika Validator & Sanitasi Regex (Email & Sandi Kuat)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "LOGIKA VALIDASI",
    "Aturan Validator: Nilai null vs Pesan Error String",
    [
        "Cara kerja fungsi validator: Menerima masukan teks yang sedang diketik '(String? value)'.",
        "Jika masukan VALID: Fungsi WAJIB mengembalikan nilai 'null' (artinya: tidak ada galat, hijau!).",
        "Jika masukan SALAH: Fungsi mengembalikan 'String' berisi pesan teguran yang otomatis tampil merah di bawah kolom.",
        "Ekspresi Reguler (Regex): Digunakan untuk menyaring pola email resmi kampus dan kekuatan sandi."
    ],
    r"""TextFormField(
  controller: _emailCtrl,
  decoration: const InputDecoration(labelText: 'Email Kampus'),
  validator: (value) {
    if (value == null || value.trim().isEmpty) {
      return 'Email tidak boleh kosong!';
    }
    // Regex pola email resmi mahasiswa kampus:
    final regex = RegExp(r'^[\w\.-]+@mhs\.kampus\.ac\.id$');
    if (!regex.hasMatch(value.trim())) {
      return 'Gunakan domain resmi: @mhs.kampus.ac.id';
    }
    return null; // LULUS VALIDASI!
  },
);""",
    filename="validator_regex.dart",
    tip="Gunakan method .trim() sebelum validasi agar spasi yang tidak sengaja tertekan tidak merusak data!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_10_validator_regex.dart"
)

# ---------------------------------------------------------------------
# Slide 11: Umpan Balik Pengguna: SnackBar & ScaffoldMessenger
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "UMPAN BALIK PENGGUNA",
    "Memberi Kabar: Menampilkan SnackBar Mengambang",
    [
        "Prinsip UX: Jangan pernah membiarkan pengguna bengong tanpa kejelasan setelah memencet tombol submit!",
        "ScaffoldMessenger.of(context).showSnackBar(): Menampilkan bilah notifikasi melayang di dasar layar ponsel.",
        "Gunakan durasi proporsional: 'Duration(seconds: 2)' agar pengguna sempat membaca tanpa merasa terganggu.",
        "Warna komunikatif: Hijau Mint untuk aksi sukses, dan Merah Coral untuk kegagalan koneksi."
    ],
    """void tampilkanSukses(BuildContext context, String nama) {
  ScaffoldMessenger.of(context).showSnackBar(
    SnackBar(
      content: Row(
        children: [
          const Icon(Icons.check_circle, color: Colors.black),
          const SizedBox(width: 10),
          Text('Selamat Datang, $nama!',
            style: const TextStyle(color: Colors.black, fontWeight: FontWeight.bold)),
        ],
      ),
      backgroundColor: const Color(0xFF4ADE80), // Hijau Mint Neo-Brutal
      behavior: SnackBarBehavior.floating,
      duration: const Duration(seconds: 2),
    ),
  );
}""",
    filename="snackbar_feedback.dart",
    tip="Pasang properti 'behavior: SnackBarBehavior.floating' agar SnackBar melayang rapi di atas bilah bawah!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_11_snackbar_feedback.dart"
)

# ---------------------------------------------------------------------
# Slide 12: Interaksi Sentuhan: GestureDetector vs InkWell
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "GESTUR & SENTUHAN",
    "Membuat Elemen Bisa Diklik: GestureDetector vs InkWell",
    [
        "GestureDetector: Pendeteksi sentuhan serba guna (tap, double tap, drag, long press), namun bekerja secara mentah TANPA animasi visual bawaan.",
        "InkWell: Pendeteksi sentuhan khas Material Design yang memancarkan animasi riak air (Ripple Splash Effect).",
        "Syarat Wajib InkWell: Harus berada di atas widget Material agar efek cipratan tintanya terlihat nyata di layar.",
        "Sangat ideal untuk membuat kartu menu kustom berdesain Neo-Brutalism!"
    ],
    """// Kartu Interaktif Berdesain Neo-Brutalism:
Material(
  color: const Color(0xFFFFE600), // Kuning Canary
  child: InkWell(
    onTap: () {
      print('Kartu pendaftaran diklik!');
    },
    splashColor: Colors.black26, // Efek riak air
    child: const Padding(
      padding: EdgeInsets.symmetric(horizontal: 24, vertical: 16),
      child: Text(
        'SUBMIT PENDAFTARAN',
        style: TextStyle(fontWeight: FontWeight.bold, fontSize: 16),
      ),
    ),
  ),
);""",
    filename="inkwell_gestures.dart",
    tip="Selalu utamakan InkWell pada tombol dan kartu agar pengguna merasakan umpan balik sentuhan yang nyata!",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_12_inkwell_gestures.dart"
)

# ---------------------------------------------------------------------
# Slide 13: Mengapa Meninggalkan Navigator 1.0? (Imperative vs Declarative)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "ARSITEKTUR NAVIGASI",
    "Mengapa Kita Butuh GoRouter? (Imperative vs Declarative)",
    [
        "Kelemahan Navigator 1.0 (Kuno): 'Navigator.push()' bekerja seperti tumpukan kartu buta (Black Box Stack). Sulit sinkron dengan URL browser web, tombol 'Back' HP sering bentrok, dan mustahil menangani Deep Linking promosi.",
        "Solusi GoRouter (Standar Industri 2026): Navigasi deklaratif berbasis URL yang terpusat ('/', '/login', '/buku/:id').",
        "Satu basis kode navigasi bekerja serentak di Android, iOS, Web, dan Desktop.",
        "Sitasi Resmi: Google Flutter Navigation Guide (docs.flutter.dev/ui/navigation)."
    ],
    """// CARA LAMA (NAVIGATOR 1.0 - IMPERATIF):
// Navigator.push(context, MaterialPageRoute(builder: (ctx) => DetailPage()));
// Kelemahan: URL web tidak sinkron, sulit Deep Linking!

// CARA MODERN 2026 (GOROUTER - DEKLARATIF):
// Rute didefinisikan terpusat berbasis alamat URL:
// /           -> Beranda
// /login      -> Halaman Masuk
// /buku/:id   -> Detail Buku Spesifik

// Berpindah semudah:
context.go('/login'); // URL otomatis sinkron di browser!""",
    filename="why_gorouter.dart",
    tip="GoRouter adalah paket navigasi deklaratif resmi yang dikelola langsung oleh tim internal Flutter!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_13_why_gorouter.dart"
)

# ---------------------------------------------------------------------
# Slide 14: Konfigurasi GoRouter & MaterialApp.router
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "SETUP ROUTER UTAMA",
    "Menghubungkan GoRouter ke Gerbang Aplikasi",
    [
        "1. Deklarasikan konfigurasi rute terpusat: 'final GoRouter appRouter = GoRouter(...)'.",
        "2. initialLocation: Alamat URL awal saat aplikasi pertama kali dibuka (biasanya '/').",
        "3. routes: Daftar pasangan alamat 'path' dengan widget tampilan 'builder'.",
        "4. MaterialApp.router: Pasang 'routerConfig: appRouter' menggantikan properti 'home:' lama.",
        "5. errorBuilder: Halaman penyelamat 404 jika pengguna membuka rute yang tidak terdaftar."
    ],
    """final appRouter = GoRouter(
  initialLocation: '/',
  routes: [
    GoRoute(path: '/', builder: (context, state) => const HomePage()),
    GoRoute(path: '/daftar', builder: (context, state) => const FormPage()),
  ],
  // Penanganan Galat Halaman 404 Ramah Awam:
  errorBuilder: (context, state) => Scaffold(
    body: Center(child: Text('404: Rute ${state.uri} Tidak Ditemukan!')),
  ),
);

void main() {
  // Hubungkan router ke aplikasi utama:
  runApp(MaterialApp.router(routerConfig: appRouter));
}""",
    filename="app_router_config.dart",
    tip="Halaman errorBuilder wajib dipasang agar aplikasi tidak mengalami layar abu-abu crash saat salah URL!",
    tag_color=COLOR_YELLOW,
    full_code_file="slide_14_app_router_config.dart"
)

# ---------------------------------------------------------------------
# Slide 15: Cara Berpindah Halaman: context.go() vs context.push()
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "METODE PINDAH RUTE",
    "Jangan Tertukar: Kapan Memakai .go() vs .push()?",
    [
        "context.go('/rute'): Mengarahkan langsung ke URL tujuan dan merombak tumpukan rute (Cocok untuk: Selesai Login -> Masuk Beranda, atau Ganti Tab Menu Bawah).",
        "context.push('/rute'): Menumpuk layar baru di atas layar saat ini (Cocok untuk: Buka Detail Transaksi atau Form Edit). Tombol panah kembali (Back) di AppBar otomatis aktif!",
        "context.pop(): Menutup layar aktif dan kembali ke layar sebelumnya."
    ],
    """// Kasus 1: Selesai Login -> Masuk Beranda (Ganti Rute):
void masukBeranda(BuildContext context) {
  context.go('/dashboard'); // Tidak bisa tekan 'Back' ke Login!
}

// Kasus 2: Klik Produk -> Buka Detail (Menumpuk Layar):
void bukaDetail(BuildContext context) {
  context.push('/detail'); // Ada tombol panah 'Back' di AppBar
}

// Kasus 3: Tombol Tutup / Selesai:
void kembali(BuildContext context) {
  context.pop(); // Menutup layar aktif saat ini
}""",
    filename="navigation_actions.dart",
    tip="Ingat rumus mudahnya: Ingin ada tombol panah 'Back' otomatis di pojok kiri atas? Gunakan context.push()!",
    tag_color=COLOR_CYAN,
    full_code_file="slide_15_navigation_actions.dart"
)

# ---------------------------------------------------------------------
# Slide 16: Kirim Data 1: Path Parameters (/buku/:id)
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "PENGIRIMAN DATA 1",
    "Mengirim ID Spesifik Menggunakan Path Parameters",
    [
        "Format URL RESTful: Gunakan tanda titik dua '/:id' pada path rute (contoh: '/buku/:id').",
        "Cara Berpindah: Panggil 'context.push(\"/buku/42\")'.",
        "Cara Membaca di Layar Tujuan: Ambil nilai ID melalui 'state.pathParameters[\"id\"]!'.",
        "Nilai yang diterima selalu bertipe String. Jika butuh angka, gunakan 'int.parse()'."
    ],
    """// 1. Daftarkan rute dengan placeholder ':id':
GoRoute(
  path: '/buku/:id',
  builder: (context, state) {
    // Ambil ID dari URL:
    final bookId = state.pathParameters['id']!;
    return DetailBukuScreen(idBuku: bookId);
  },
);

// 2. Berpindah dari layar katalog buku:
void pilihBuku(BuildContext context, String id) {
  context.push('/buku/$id'); // Menuju '/buku/101'
}""",
    filename="path_parameters.dart",
    tip="Path Parameters sangat ideal untuk identifier unik seperti NIM, ID Produk, atau Nomor Transaksi!",
    tag_color=COLOR_PURPLE,
    full_code_file="slide_16_path_parameters.dart"
)

# ---------------------------------------------------------------------
# Slide 17: Kirim Data 2: Query Parameters & Extra Object
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "PENGIRIMAN DATA 2",
    "Kueri Filter (?q=) & Mengirim Objek Kompleks (Extra)",
    [
        "Query Parameters: Menyematkan filter pencarian opsional pada URL (misal: '/cari?keyword=flutter'). Dibaca lewat 'state.uri.queryParameters[\"keyword\"]'.",
        "Extra Object (extra): Mengirim objek class Dart utuh (misal: objek Mahasiswa dengan 10 atribut) tanpa repot dipecah menjadi string URL!",
        "Dibaca di halaman tujuan dengan: 'state.extra as NamaModel'."
    ],
    """// Mengirim objek Mahasiswa secara utuh:
context.push(
  '/profil',
  extra: Mahasiswa(nim: '23001', nama: 'Adit', ipk: 3.85),
);

// Membaca objek di definisi rute GoRouter:
GoRoute(
  path: '/profil',
  builder: (context, state) {
    final mhs = state.extra as Mahasiswa;
    return ProfilScreen(mahasiswa: mhs);
  },
);""",
    filename="query_and_extra.dart",
    tip="Gunakan 'extra' saat Anda ingin mengirim objek data kompleks tanpa mengotori alamat URL browser!",
    tag_color=COLOR_MINT,
    full_code_file="slide_17_query_and_extra.dart"
)

# ---------------------------------------------------------------------
# Slide 18: Rute Bersarang (Nested Routes) & StatefulShellRoute
# ---------------------------------------------------------------------
deck4.add_concept_with_code(
    "RUTE BERSARANG & SHELL",
    "Rute Sub-Halaman & Navigasi Tab Bottom Navigation",
    [
        "Rute Bersarang (Nested Routes): Menempatkan sub-rute di dalam parent GoRoute ('/katalog' -> 'detail').",
        "URL otomatis menjadi hierarkis: '/katalog/detail'. Sangat teratur dan bersih.",
        "StatefulShellRoute.indexedStack(): Menjaga status tumpukan memori masing-masing tab Bottom Navigation Bar agar tidak ter-reset saat berpindah.",
        "Standar emas arsitektur navigasi aplikasi mobile skala besar 2026."
    ],
    """// 1. Rute Bersarang (Sub-Routes):
GoRoute(
  path: '/katalog',
  builder: (ctx, state) => const KatalogScreen(),
  routes: [
    GoRoute(
      path: 'detail/:id', // Menjadi '/katalog/detail/42'
      builder: (ctx, state) => DetailScreen(id: state.pathParameters['id']!),
    ),
  ],
);

// 2. Tab Bawah Modern: StatefulShellRoute.indexedStack(...)""",
    filename="nested_and_shell_routes.dart",
    tip="Gunakan nested routes agar struktur alamat URL aplikasi Anda mencerminkan hierarki menu secara akurat!",
    tag_color=COLOR_CORAL,
    full_code_file="slide_18_nested_and_shell_routes.dart"
)

# ---------------------------------------------------------------------
# Slide 19: Lab Quest Mandiri Pertemuan 04 (60 Menit)
# ---------------------------------------------------------------------
deck4.add_lab_quest(
    4,
    "Portal Registrasi Mahasiswa & Detail KTM Ber-GoRouter",
    60,
    [
        "Buka terminal terintegrasi di VS Code dan pasang pustaka: 'flutter pub add go_router'.",
        "Buat Halaman Formulir Registrasi (/daftar) yang dibungkus widget 'Form' dan 'GlobalKey<FormState>'.",
        "Sediakan kolom: Nama Lengkap, Email Kampus (@mhs.kampus.ac.id), Sandi dengan toggle intip mata, dan Dropdown Program Studi.",
        "Gunakan 'TextEditingController' dan pastikan seluruh controller dibersihkan pada method 'dispose()'.",
        "Tutup keyboard virtual dengan 'FocusScope.of(context).unfocus()' saat tombol submit ditekan.",
        "Ketika submit sukses: Tampilkan SnackBar hijau dan navigasi ke '/ktm/:nim' menggunakan 'context.push()' dengan membawa parameter NIM dan extra objek Mahasiswa.",
        "Halaman KTM (/ktm/:nim): Menampilkan kartu KTM digital bergaya Neo-Brutalism dengan tombol 'Kembali' menggunakan 'context.pop()'.",
        "Tantangan Tambahan: Tambahkan halaman 'errorBuilder' 404 jika ada URL yang sengaja diketik salah!"
    ],
    "Alur pendaftaran berjalan mulus dari form ke kartu KTM, tervalidasi regex, bebas kebocoran memori (dispose aktif), keyboard menutup anggun, dan navigasi GoRouter bekerja sempurna di browser Chrome atau Smartphone!"
)

# ---------------------------------------------------------------------
# Simpan berkas PPTX Pertemuan 04
# ---------------------------------------------------------------------
output_path = os.path.join(output_dir, "PERTEMUAN_04_Form_Handling_dan_GoRouter.pptx")
deck4.save(output_path)
